# -*- coding: utf-8 -*-
"""Таблица площадок из SMI-PLAN.md: каждая строка таблиц изданий -> строка xlsx."""
import re
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment

CVET = {'открыт': 'E2F0D9', 'жив, режет': 'FFF2CC', 'не проверено': 'F4CCCC'}

wb = Workbook()
ws = wb.active
ws.title = 'Издания'
ws.append(['Направление', 'Издание', 'Статус', 'Контакт', 'Что даёт',
           'Дата обращения', 'Тема', 'Вышло', 'Ссылка', 'dofollow', 'Заявки'])
razdel = ''
for s in open('SMI-PLAN.md', encoding='utf-8'):
    if s.startswith('### '):
        razdel = s[4:].strip()
    if not razdel or not s.startswith('|') or s.startswith('|---') or s.startswith('| Издание'):
        continue
    k = [x.strip().replace('**', '') for x in s.strip().strip('|').split('|')]
    if len(k) == 3:
        k = [k[0], k[1], '', k[2]]
    ws.append([razdel] + k)
    st = next((v for v in CVET if k[1].startswith(v)), None)
    if st:
        ws.cell(ws.max_row, 3).fill = PatternFill('solid', fgColor=CVET[st])
for c in ws[1]:
    c.font = Font(bold=True)
for kol, w in zip('ABCDEFGHIJK', (34, 40, 22, 30, 60, 14, 30, 8, 30, 9, 8)):
    ws.column_dimensions[kol].width = w
for r in ws.iter_rows(min_row=2):
    for c in r:
        c.alignment = Alignment(wrap_text=True, vertical='top')
ws.freeze_panes = 'A2'
wb.save('smi-ploshchadki.xlsx')
print('строк:', ws.max_row - 1)
