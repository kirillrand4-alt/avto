# -*- coding: utf-8 -*-
"""Эксель базы под Мейер: шесть сегментов, источник у каждой строки.

ОТКУДА ВЗЯТО, по полям:
  наша база enrich.db / companies ........ название, регион, сайт, вид деятельности,
                                           почта, метка division (meyer / kc)
  наша база enrich.db / phone_contacts ... телефон, имя, роль, ИСТОЧНИК и ССЫЛКА.
                                           778 тысяч строк, у 98 % есть source_url
  наша база enrich.db / people ........... названные люди с должностями
  справочник obzvon-index.db ............. ОКВЭД основной и дополнительные, выручка,
                                           год отчётности, численность, директор ЕГРЮЛ

Главное число этого файла - не 43 559 компаний, а вид контакта. Владелец просил не
секретарей, значит надо видеть, где секретарь. Поэтому колонка «Вид контакта» стоит
третьей, а не спрятана в конце.
"""
import collections
import csv
import io
import os

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

KAT = os.path.dirname(os.path.abspath(__file__))
VYHOD = os.path.join(KAT, 'MEYER-BAZA.xlsx')
SHRIFT = 'Arial'

komp = list(csv.DictReader(io.open(os.path.join(KAT, 'MEYER-BAZA-KOMPANII.csv'),
                                   encoding='utf-8-sig'), delimiter=';'))
kont = list(csv.DictReader(io.open(os.path.join(KAT, 'MEYER-BAZA-KONTAKTY.csv'),
                                   encoding='utf-8-sig'), delimiter=';'))
print('компаний %d, контактов %d' % (len(komp), len(kont)))


def ch(x, k):
    try:
        return float(str(x.get(k) or '').replace(',', '.'))
    except ValueError:
        return None


# сортировка: сначала те, у кого есть живой человек, потом по выручке
RANG_VIDA = {'личный контакт с именем': 0, 'роль ЛПР без имени': 1, 'первое лицо': 2,
             'директор по ЕГРЮЛ (имя есть, прямого номера нет)': 3,
             'приёмная или общий номер': 4, 'номер без роли': 5, 'контакта нет': 6}
komp.sort(key=lambda x: (RANG_VIDA.get(x['vid_kontakta'], 9), -(ch(x, 'vyruchka_rub') or 0)))

LUCHSHIE = [x for x in kont if x['fio'].strip() or x['rang'] in ('1', '2', '3', '4')]
OBSHIE = [x for x in kont if x['rang'] == '5']
LUCHSHIE.sort(key=lambda x: (x['rang'], x['nazvanie']))
print('лучших контактов %d, приёмных и общих %d' % (len(LUCHSHIE), len(OBSHIE)))

vid = collections.Counter(x['vid_kontakta'] for x in komp)
seg = collections.Counter()
for x in komp:
    for s in x['segmenty'].split(' | '):
        if s.strip():
            seg[s.strip()] += 1

wb = Workbook()
ZAG = Font(name=SHRIFT, size=12, bold=True)
OB = Font(name=SHRIFT, size=10)
ws = wb.active
ws.title = 'Как читать'
OPIS = [
    ('База под Мейер: шесть сегментов, источник у каждой строки', True),
    ('', False),
    ('ОТКУДА ВЗЯТО, по полям', True),
    ('   наша база enrich.db / companies ..... название, регион, сайт, вид деятельности,', False),
    ('                                        почта, метка division (meyer или kc)', False),
    ('   наша база enrich.db / phone_contacts  телефон, имя, роль, ИСТОЧНИК и ССЫЛКА.', False),
    ('                                        778 тысяч строк, у 98 % есть ссылка', False),
    ('   наша база enrich.db / people ........ названные люди с должностями', False),
    ('   справочник obzvon-index.db .......... ОКВЭД основной и дополнительные, выручка,', False),
    ('                                        год отчётности, численность, директор ЕГРЮЛ', False),
    ('   У каждой строки стоит колонка «Ссылка на источник контакта», а для реквизитов -', False),
    ('   «Ссылка ЕГРЮЛ» для проверки по ИНН за минуту.', False),
    ('', False),
    ('СЕГМЕНТЫ ровно по вашему списку ОКВЭД', True),
] + [('   %-50s %6d' % (k, n), False) for k, n in sorted(seg.items())] + [
    ('   Компания может попасть в несколько сегментов - тогда в колонке «Сегменты» их', False),
    ('   несколько через черту. Сумма по сегментам больше числа компаний, и это не ошибка.', False),
    ('   Всего разных компаний: %d, из них действующих %d.'
     % (len(komp), len([x for x in komp if 'ействующ' in x['status']])), False),
    ('', False),
    ('ГЛАВНОЕ ЧИСЛО ФАЙЛА - НЕ 43 559 КОМПАНИЙ, А ВИД КОНТАКТА', True),
] + [('   %-52s %6d' % (k, n), False) for k, n in vid.most_common()] + [
    ('', False),
    ('   Вы просили не секретарей, а прямые контакты ЛПР. Честно: таких в собранном мало.', False),
    ('   Личный контакт с именем есть у %d компаний из %d. Ещё у %d стоит роль ЛПР без'
     % (vid.get('личный контакт с именем', 0), len(komp), vid.get('роль ЛПР без имени', 0)), False),
    ('   имени, у %d - первое лицо. Самая массовая строка - «директор по ЕГРЮЛ»: имя есть,'
     % vid.get('первое лицо', 0), False),
    ('   прямого номера нет, звонить придётся через приёмную.', False),
    ('   Это не свойство сегмента, это состояние базы: во всей таблице phone_contacts имя', False),
    ('   человека заполнено у 3 068 строк из 777 736, то есть у 546 ИНН на всю базу.', False),
    ('', False),
    ('ЛИСТЫ', True),
    ('   «Компании»          %6d строк. Сверху те, у кого есть живой человек, дальше по' % len(komp), False),
    ('                       убыванию выручки. Фильтры включены.', False),
    ('   «Лучшие контакты»   %6d строк: имя есть ИЛИ роль относится к ЛПР (техническая,' % len(LUCHSHIE), False),
    ('                       качество, закупки, первое лицо, производство). Со ссылками.', False),
    ('   «Приёмные и общие»  %6d строк: приёмная, секретарь, продажи, бухгалтерия.' % len(OBSHIE), False),
    ('                       Не выброшены: через них выходят на человека.', False),
    ('   Полные выгрузки без обрезки - на дропе: MEYER-BAZA-KOMPANII.csv (52 МБ) и', False),
    ('   MEYER-BAZA-KONTAKTY.csv (50 МБ, 212 958 строк контактов).', False),
    ('', False),
    ('ЧЕГО В ФАЙЛЕ НЕТ, сказано прямо', True),
    ('   • БЕЛАРУСИ НЕТ. Вы просили добавить белорусов в ягодный сегмент - в собранном их', False),
    ('     нет вовсе: по слову «Беларус» 3 строки, «Минск» 84, и это российские компании', False),
    ('     с минским адресом в тексте. Белорусов надо добывать отдельно, это новый сбор.', False),
    ('   • ЭКСПОРТЁРЫ - это КАНДИДАТЫ. ОКВЭД экспорта не существует, а признака ВЭД или', False),
    ('     таможенных операций в нашей базе нет ни одного. Сегмент собран по оптовой', False),
    ('     торговле зерном и продуктами переработки (46.21, 46.31-46.38, 46.17).', False),
    ('     Это гипотеза по отрасли, а не подтверждённый экспорт, и так подписано.', False),
    ('   • выручка есть у %d компаний из %d. Она НЕ использовалась как фильтр - вы просили'
     % (len([x for x in komp if x['vyruchka']]), len(komp)), False),
    ('     мелких не выбрасывать. Колонка «Выручка, руб» числовая, по ней можно сортировать.', False),
    ('   • сайт есть у %d, почта у %d, рабочий телефон у %d.'
     % (len([x for x in komp if x['sayt']]), len([x for x in komp if x['email']]),
        len([x for x in komp if x['rabochiy_telefon']])), False),
    ('', False),
    ('КОНТРОЛИ', True),
    ('   • выдуманный ИНН 9999999999 в выборке даёт 0 строк.', False),
    ('   • из 43 559 компаний ВСЕ 43 559 нашлись в нашей базе companies - то есть сегменты', False),
    ('     собраны по нашей базе, а не по чужому справочнику.', False),
    ('   • строк контактов со ссылкой на источник: 207 159 из 212 958 (97 %).', False),
    ('   • в первом заходе я join-илась только к таблице people и получила ОДНОГО', False),
    ('     названного человека на 43 559 компаний. Ошибка была в запросе, а не в данных:', False),
    ('     имя и ссылка лежат ещё и в phone_contacts. Числа выше - после исправления.', False),
]
for i, (t, zh) in enumerate(OPIS, start=1):
    c = ws.cell(row=i, column=1, value=t)
    c.font = ZAG if zh else OB
ws.column_dimensions['A'].width = 100

SH_K = [('Сегменты', 'segmenty', 46), ('Название', 'nazvanie', 42),
        ('Вид контакта', 'vid_kontakta', 30), ('ИНН', 'inn', 13),
        ('Регион', 'region', 22), ('Город и адрес', 'adres', 34),
        ('Сайт', 'sayt', 28), ('Основной ОКВЭД', 'okved_osnovnoy', 46),
        ('Дополнительные ОКВЭД', 'okved_dopolnitelnye', 40),
        ('Описание деятельности', 'opisanie_deyatelnosti', 40),
        ('Выручка', 'vyruchka', 15), ('Выручка, руб', 'vyruchka_rub', 16),
        ('Год отчётности', 'god_otchetnosti', 12), ('Сотрудников', 'sotrudnikov', 11),
        ('ФИО контакта', 'fio_kontakta', 28), ('Должность или роль', 'dolzhnost_ili_rol', 30),
        ('Прямой телефон', 'pryamoy_telefon', 18), ('Рабочий телефон', 'rabochiy_telefon', 18),
        ('E-mail', 'email', 26), ('Источник контакта', 'istochnik_kontakta', 24),
        ('Ссылка на источник контакта', 'ssylka_na_istochnik_kontakta', 44),
        ('Контактов всего', 'kontaktov_vsego', 11), ('Из них с именем', 'kontaktov_s_imenem', 11),
        ('Из них ЛПР', 'kontaktov_rang_1_2', 10),
        ('Другие контакты', 'drugie_kontakty', 44),
        ('Метка нашей базы', 'razmetka_nashey_bazy', 13),
        ('Статус', 'status', 20), ('Ссылка ЕГРЮЛ', 'ssylka_egryul', 40)]
w1 = wb.create_sheet('Компании')
w1.append([a for a, _, _ in SH_K])
for x in komp:
    r = []
    for _, k, _w in SH_K:
        v = x.get(k, '')
        if k == 'vyruchka_rub':
            v = ch(x, k)
        elif k in ('kontaktov_vsego', 'kontaktov_s_imenem', 'kontaktov_rang_1_2'):
            v = int(v or 0)
        elif k in ('okved_dopolnitelnye', 'drugie_kontakty', 'adres'):
            v = str(v)[:200]
        r.append(v)
    w1.append(r)

SH_C = [('ИНН', 'inn', 13), ('Название', 'nazvanie', 42), ('Сегменты', 'segmenty', 44),
        ('ФИО', 'fio', 30), ('Роль как в источнике', 'rol_kak_v_istochnike', 34),
        ('Группа ЛПР', 'gruppa_lpr', 26), ('Ранг', 'rang', 7),
        ('Телефон', 'telefon', 20), ('Источник', 'istochnik', 24),
        ('Ссылка на источник', 'ssylka_na_istochnik', 60)]


def list_kont(imya, gr):
    w = wb.create_sheet(imya)
    w.append([a for a, _, _ in SH_C])
    for x in gr:
        w.append([x.get(k, '') for _, k, _w in SH_C])
    return w


w2 = list_kont('Лучшие контакты', LUCHSHIE)
w3 = list_kont('Приёмные и общие', OBSHIE)

SH_F = Font(name=SHRIFT, size=10, bold=True, color='FFFFFF')
SH_FILL = PatternFill('solid', fgColor='2F5496')
for w, shir in ((w1, [c for _, _, c in SH_K]), (w2, [c for _, _, c in SH_C]),
                (w3, [c for _, _, c in SH_C])):
    for c in w[1]:
        c.font = SH_F
        c.fill = SH_FILL
        c.alignment = Alignment(vertical='center', wrap_text=True)
    w.freeze_panes = 'A2'
    w.auto_filter.ref = w.dimensions
    w.row_dimensions[1].height = 32
    for i, s in enumerate(shir, start=1):
        w.column_dimensions[get_column_letter(i)].width = s
    for row in w.iter_rows(min_row=2):
        for c in row:
            c.font = Font(name=SHRIFT, size=10)
            c.alignment = Alignment(vertical='top')

wb.save(VYHOD)
print('файл: %s (%.1f МБ)' % (VYHOD, os.path.getsize(VYHOD) / 1048576.0))
print('  Компании %d · Лучшие контакты %d · Приёмные и общие %d'
      % (len(komp), len(LUCHSHIE), len(OBSHIE)))
