# -*- coding: utf-8 -*-
"""Только «наша тема» отдельным файлом: 294 объекта, где компрессорная — предмет проекта.

Владелец попросил этот лист отдельно. Здесь нет ни АГНКС, ни медицинского кислорода,
ни отсева — они остаются в полном файле EGRZ-NASHA-TEMA.xlsx. Контакты и поводы тоже
урезаны: только по ИНН, которые встречаются на этом листе.
"""
import csv
import io
import json
import os
import re

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

KAT = os.path.dirname(os.path.abspath(__file__))
VHOD = os.path.join(KAT, 'EGRZ-NASHA-TEMA-2.jsonl')
VYHOD = os.path.join(KAT, 'EGRZ-NASHA-TEMA-TOLKO.xlsx')
SHRIFT = 'Arial'
TEH = re.compile(r'инженер|энергетик|механик|технолог|техдирект|техн|производств|цех|'
                 r'КИПиА|АСУ|снабжен|закупк|заказчик|капитальн|проект|эксплуатац', re.I)

stroki = [json.loads(s) for s in io.open(VHOD, encoding='utf-8') if s.strip()]
for s in stroki:
    for p in ('zastroyshchik_inn', 'proektirovshchik_inn', 'tehzakazchik_inn'):
        if set(s.get(p) or 'x') == {'0'}:
            s[p] = ''
NASHA = sorted([s for s in stroki if s['klass'] == 'НАША ТЕМА'],
               key=lambda z: z.get('data') or '', reverse=True)

NOVIZNA = {r['inn']: r['novoe_dlya_bazy'] for r in csv.DictReader(
    io.open(os.path.join(KAT, '3s_NT-NOVIZNA.csv'), encoding='utf-8-sig'), delimiter=';')}

# ИНН только этого листа: контакты и поводы урезаются по нему, иначе в файле окажутся
# телефоны заправок и больниц, которых на листе нет.
NASHI_INN = set()
for s in NASHA:
    for p in ('zastroyshchik_inn', 'proektirovshchik_inn', 'tehzakazchik_inn'):
        if s.get(p):
            NASHI_INN.add(s[p])

kont = [k for k in csv.DictReader(io.open(os.path.join(KAT, '3s_NT-KONTAKTY.csv'),
                                          encoding='utf-8-sig'), delimiter=';')
        if k['inn'] in NASHI_INN]
povod = [p for p in csv.DictReader(io.open(os.path.join(KAT, '3s_NT-POVODY.csv'),
                                           encoding='utf-8-sig'), delimiter=';')
         if p['inn'] in NASHI_INN]
for k in kont:
    k['tehnicheskaya'] = 'да' if TEH.search((k.get('chelovek_dolzhnost') or '') + ' '
                                            + (k.get('rol_kanon') or '')) else ''
po_tel, po_teh, po_pov = {}, {}, {}
for k in kont:
    i = k['inn']
    if (k.get('telefon') or '').strip():
        po_tel[i] = po_tel.get(i, 0) + 1
    if k['tehnicheskaya']:
        po_teh[i] = po_teh.get(i, 0) + 1
for p in povod:
    po_pov[p['inn']] = po_pov.get(p['inn'], 0) + 1


def raznyh(pole):
    return len({s[pole] for s in NASHA if s.get(pole)})


def novyh(pole):
    return len({s[pole] for s in NASHA if s.get(pole) and NOVIZNA.get(s[pole]) == 'да'})


wb = Workbook()
ZAG = Font(name=SHRIFT, size=12, bold=True)
OB = Font(name=SHRIFT, size=10)
ws = wb.active
ws.title = 'Как читать'
OPIS = [
    ('Наша тема: кто строит компрессорную и кто ему это проектирует', True),
    ('', False),
    ('Это ОДИН лист из полного файла EGRZ-NASHA-TEMA.xlsx, вынесенный отдельно.', False),
    ('   Здесь только то, где компрессорная машина является предметом проекта.', False),
    ('   АГНКС (709 строк), медицинский кислород (88), смежное производство (77) и весь', False),
    ('   отсев (233 строки с причиной у каждой) остались в полном файле - не выброшены.', False),
    ('', False),
    ('ИСТОЧНИК. ЕГРЗ - единый реестр заключений экспертизы проектной документации,', False),
    ('   open-api.egrz.ru, ключ не нужен. В одной записи стоят ОБЕ стороны с ИНН:', False),
    ('   застройщик (будущий эксплуатант) и ПРОЕКТИРОВЩИК, который на стадии проекта и', False),
    ('   выбирает оборудование. Это событие ДО стройки, а не после неё.', False),
    ('', False),
    ('ЧИСЛА ЭТОГО ФАЙЛА', True),
    ('   объектов .............................. %d' % len(NASHA), False),
    ('   разных застройщиков ................... %d, из них НОВЫХ для нашей базы %d'
     % (raznyh('zastroyshchik_inn'), novyh('zastroyshchik_inn')), False),
    ('   разных проектировщиков ................ %d, из них НОВЫХ для нашей базы %d'
     % (raznyh('proektirovshchik_inn'), novyh('proektirovshchik_inn')), False),
    ('   объектов, где у застройщика есть телефон .... %d'
     % len([s for s in NASHA if po_tel.get(s.get('zastroyshchik_inn', ''))]), False),
    ('   из них телефон при технической роли ......... %d'
     % len([s for s in NASHA if po_teh.get(s.get('zastroyshchik_inn', ''))]), False),
    ('   объектов, по застройщику которых есть повод .. %d'
     % len([s for s in NASHA if po_pov.get(s.get('zastroyshchik_inn', ''))]), False),
    ('   Проектировщики новы почти поголовно: в базе копились эксплуатанты, а машину на', False),
    ('   стадии проекта выбирает проектировщик. Колонки «новый для базы» стоят в строках.', False),
    ('', False),
    ('ЛИСТЫ', True),
    ('   «Наша тема»  %d объектов, свежие сверху (сортировка по дате заключения)' % len(NASHA), False),
    ('   «Контакты»   %d строк - что уже есть в нашей базе по ИНН С ЭТОГО ЛИСТА' % len(kont), False),
    ('   «Поводы»     %d строк - события по тем же ИНН' % len(povod), False),
    ('', False),
    ('ЧЕГО ЗДЕСЬ НЕТ, сказано прямо:', True),
    ('   • людей в самом ЕГРЗ нет вовсе - ни ФИО, ни телефонов. Контакты подтянуты из нашей', False),
    ('     базы и только для тех ИНН, которые в ней уже были.', False),
    ('   • марки машины в ЕГРЗ практически нет: какой именно компрессор стоит - не отсюда.', False),
    ('   • у 13 объектов застройщик не назван вовсе - так в самом реестре. Проектировщик', False),
    ('     у таких обычно есть, и звонить надо ему.', False),
    ('   • это ЗАКЛЮЧЕНИЯ экспертизы. Стройка могла не начаться, а могла и закончиться -', False),
    ('     смотрите колонку «Дата заключения».', False),
    ('   • класс ставит правило, а не человек. 18 строк, где правило не справилось, лежат', False),
    ('     на листе «Сомнительно» полного файла - сюда они не попали.', False),
    ('', False),
    ('КОНТРОЛИ, без которых числам верить нельзя:', True),
    ('   • выдуманное слово «нипрятозаумень» даёт в реестре 0 записей - прибор различает.', False),
    ('   • выдуманный ИНН 9999999999 в нашей базе даёт 0 контактов и 0 поводов.', False),
    ('   • сверка новизны: выдуманный ИНН обязан оказаться новым (оказался), взятый из', False),
    ('     самой базы - не новым (не оказался).', False),
    ('   • поиск ТОЛЬКО через tolower(): contains() в ЕГРЗ регистрозависим, наивный поиск', False),
    ('     теряет четверть («компрессорная» 208, «Компрессорная» 65, через tolower 274).', False),
    ('   • отбор проверен глазами: на 30 случайных строках этого листа и на всех строках', False),
    ('     прошлой выгрузки, ушедших в отсев. Так нашлись и подделки (нагнетательная', False),
    ('     СКВАЖИНА, переделка компрессорной под офис, изготовитель компрессоров, топоним),', False),
    ('     и потери («НЛМК. Кислородный цех», «азотно-воздушная станция», «ЭВД-4 с', False),
    ('     установкой компрессора») - последние в этот лист возвращены.', False),
    ('', False),
    ('Пять последних колонок посчитаны на дату выгрузки и лежат числами, а не формулами:', False),
    ('   это снимок, а не модель. Правите данные - счётчики сами не изменятся.', False),
]
for i, (t, zh) in enumerate(OPIS, start=1):
    c = ws.cell(row=i, column=1, value=t)
    c.font = ZAG if zh else OB
ws.column_dimensions['A'].width = 100

SHAPKA = ['Дата заключения', 'Регион', 'Объект', 'Адрес объекта', 'Вывод экспертизы',
          'Вид документа', 'Застройщик', 'ИНН застройщика', 'Застройщик новый для базы',
          'Адрес застройщика', 'Проектировщик', 'ИНН проектировщика',
          'Проектировщик новый для базы', 'Адрес проектировщика', 'Технический заказчик',
          'ИНН техзаказчика', 'Найден по слову', 'Номер заключения', 'Ссылка на карточку',
          'Телефонов у застройщика', 'Из них тех.роль', 'Поводов у застройщика',
          'Телефонов у проектировщика', 'Из них тех.роль']
SHIR = [14, 22, 58, 32, 20, 20, 38, 14, 13, 30, 38, 16, 13, 30, 30, 15, 20, 20, 42,
        12, 12, 12, 13, 12]
w1 = wb.create_sheet('Наша тема')
w1.append(SHAPKA)
for s in NASHA:
    zi, pi = s.get('zastroyshchik_inn', ''), s.get('proektirovshchik_inn', '')
    w1.append([s.get('data', ''), s.get('region', ''), s.get('obekt', ''),
               s.get('adres_obekta', ''), s.get('vyvod', ''), s.get('vid_dokumenta', ''),
               s.get('zastroyshchik', ''), zi, NOVIZNA.get(zi, ''),
               s.get('zastroyshchik_adres', ''),
               s.get('proektirovshchik', ''), pi, NOVIZNA.get(pi, ''),
               s.get('proektirovshchik_adres', ''),
               s.get('tehzakazchik', ''), s.get('tehzakazchik_inn', ''),
               s.get('nayden_po', ''), s.get('nomer_zaklyucheniya', ''), s.get('ssylka', ''),
               po_tel.get(zi, 0), po_teh.get(zi, 0), po_pov.get(zi, 0),
               po_tel.get(pi, 0), po_teh.get(pi, 0)])

wk = wb.create_sheet('Контакты')
wk.append(['ИНН', 'Должность как в источнике', 'Роль (канон)', 'Телефон', 'Почта',
           'Источник', 'Ссылка', 'Откуда взято', 'Тех.роль'])
for k in kont:
    wk.append([k['inn'], k.get('chelovek_dolzhnost', ''), k.get('rol_kanon', ''),
               k.get('telefon', ''), k.get('pochta', ''), k.get('istochnik', ''),
               k.get('ssylka', ''), k.get('otkuda', ''), k.get('tehnicheskaya', '')])
wp = wb.create_sheet('Поводы')
wp.append(['ИНН', 'Тип события', 'Что происходит', 'Сумма', 'Источник', 'Ссылка', 'Когда'])
for p in povod:
    wp.append([p['inn'], p.get('tip_sobytiya', ''), p.get('chto', ''), p.get('summa', ''),
               p.get('istochnik', ''), p.get('ssylka', ''), p.get('kogda', '')])

SH_F = Font(name=SHRIFT, size=10, bold=True, color='FFFFFF')
SH_FILL = PatternFill('solid', fgColor='2F5496')
SHIR_DOP = {'Контакты': [14, 38, 18, 18, 28, 20, 44, 16, 10],
            'Поводы': [14, 22, 70, 18, 22, 44, 14]}
for w in (w1, wk, wp):
    for c in w[1]:
        c.font = SH_F
        c.fill = SH_FILL
        c.alignment = Alignment(vertical='center', wrap_text=True)
    w.freeze_panes = 'A2'
    w.auto_filter.ref = w.dimensions
    w.row_dimensions[1].height = 30
    for i, sh in enumerate(SHIR_DOP.get(w.title, SHIR), start=1):
        w.column_dimensions[get_column_letter(i)].width = sh
    for row in w.iter_rows(min_row=2):
        for c in row:
            c.font = Font(name=SHRIFT, size=10)
            c.alignment = Alignment(vertical='top')

wb.save(VYHOD)
print('файл: %s' % VYHOD)
print('  объектов %d · застройщиков %d (новых %d) · проектировщиков %d (новых %d)'
      % (len(NASHA), raznyh('zastroyshchik_inn'), novyh('zastroyshchik_inn'),
         raznyh('proektirovshchik_inn'), novyh('proektirovshchik_inn')))
print('  ИНН на листе %d · контактов %d (с телефоном %d, тех.роль %d) · поводов %d'
      % (len(NASHI_INN), len(kont), len([k for k in kont if k['telefon'].strip()]),
         len([k for k in kont if k['tehnicheskaya']]), len(povod)))
