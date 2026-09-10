# -*- coding: utf-8 -*-
"""Разбор годовых выгрузок реестра ЭПБ Ростехнадзора (файлы Excel с сайта территориального управления).
Формат Западно-Сибирского управления: колонка «Субъект РФ», ИНН эксплуатирующей организации, номер и класс ОПО,
срок дальнейшей безопасной эксплуатации, текст заключения с маркой машины. Листы бывают понедельные.
Берём строки Алтайского края, где в заключении названа наша машина, и делаем карточки доказательств.
Запуск локальный (файлы приходят от владельца): python3 ak_rtn_reestr.py <файл.xlsx> [ещё файлы] -> kartochki-rtn.json"""
import sys, re, json, os
from openpyxl import load_workbook

MASH = re.compile(r'компрессор|воздуходув|турбовоздуходув|нагнетател|ресивер|воздухосборник|осушител\w*\s*(?:сжатого\s*)?воздух'
                  r'|воздухоразделительн|\bВРУ\b|азотн\w+\s+(?:станц|установ)|кислородн\w+\s+(?:станц|установ)|генератор\w*\s+(?:азота|кислорода)', re.I)
VIDY = [('генератор кислорода', r'кислородн\w*\s*(станц|генератор|установ)|генератор\w*\s*кислород'),
        ('генератор азота', r'азотн\w*\s*(станц|генератор|установ)|генератор\w*\s*азот'),
        ('ВРУ', r'воздухоразделительн|\bВРУ\b'), ('воздуходувка', r'воздуходув|турбовоздуходув'),
        ('нагнетатель', r'нагнетател'), ('осушитель', r'осушител'),
        ('компрессорная станция', r'компрессорн\w*\s*(станци|установк|цех)|здани\w*\s+компрессорн'),
        ('ресивер', r'ресивер|воздухосборник'), ('компрессор', r'компрессор')]
VIDY = [(v, re.compile(r, re.I)) for v, r in VIDY]
def vid(s):
    for v, rx in VIDY:
        if rx.search(s or ''): return v
    return 'компрессор'

def kolonki(vals):
    """Шапка бывает с разным числом колонок - ищем по названиям, а не по номерам."""
    i = {}
    for n, v in enumerate(vals):
        v = v.strip()
        if v.startswith('Субъект РФ'): i['subekt'] = n
        elif v.startswith('Регистрационный'): i['reg'] = n
        elif v.startswith('Дата регистрации'): i['data'] = n
        elif v.startswith('Срок дальнейшей'): i['srok'] = n
        elif v.startswith('Вывод заключения'): i['vyvod'] = n
        elif v.startswith('Наименование ЗЭПБ'): i['zepb'] = n
        elif v.startswith('Номер ОПО'): i['opo_nom'] = n
        elif v.startswith('Наименование ОПО'): i['opo_imya'] = n
        elif v.startswith('Класс опасности'): i['klass'] = n
        elif v.startswith('Наименование эксплуатирующей'): i['imya'] = n
        elif v.startswith('ИНН эксп'): i['inn'] = n
    return i if {'subekt', 'zepb', 'inn'} <= set(i) else None

kartochki = []; svod = {'strok': 0, 'kray': 0, 'mashin': 0, 'opo': 0, 'listov': 0}
for put in sys.argv[1:]:
    wb = load_workbook(put, read_only=True)
    for sh in wb.sheetnames:
        ws = wb[sh]; idx = None; svod['listov'] += 1
        for r in ws.iter_rows(values_only=True):
            vals = [str(x) if x is not None else '' for x in r]
            if idx is None:
                idx = kolonki(vals)
                continue
            if not any(vals): continue
            svod['strok'] += 1
            def p(k):
                n = idx.get(k)
                return vals[n].strip() if n is not None and n < len(vals) else ''
            if 'Алтайский край' not in p('subekt'): continue
            svod['kray'] += 1
            zepb = re.sub(r'\s+', ' ', p('zepb'))
            opo_imya = re.sub(r'\s+', ' ', p('opo_imya'))
            # машина названа в самом заключении - это ЭПБ; названа только в имени ОПО
            # («Площадка компрессорной станции») - это факт вида ОПО, тоже доказательство, но послабее
            v_zepb = bool(MASH.search(zepb)); v_opo = bool(MASH.search(opo_imya))
            if not (v_zepb or v_opo): continue
            svod['mashin' if v_zepb else 'opo'] += 1
            inn = re.sub(r'\D', '', p('inn'))
            if len(inn) not in (10, 12): continue
            citata = (zepb if v_zepb else opo_imya)[:450]
            hvost = ' | '.join(x for x in (opo_imya if v_zepb else '', p('klass'), p('vyvod')[:60]) if x)
            kartochki.append({'inn': inn, 'imya': p('imya'), 'vid_fakta': 'ЭПБ' if v_zepb else 'ОПО',
                              'tip': vid(zepb if v_zepb else opo_imya),
                              'data': p('data'), 'srok_do': p('srok') if v_zepb else '', 'reg': p('reg'),
                              'opo': p('opo_nom'), 'klass': p('klass'),
                              'citata': (citata + (' | ' + hvost if hvost else ''))[:600],
                              'fayl': os.path.basename(put)})
    wb.close()
out = os.path.join(os.path.dirname(os.path.abspath(__file__)), 'kartochki-rtn.json')
json.dump(kartochki, open(out, 'w', encoding='utf-8'), ensure_ascii=False)
print('листов', svod['listov'], '| строк', svod['strok'], '| край', svod['kray'], '| машина в заключении', svod['mashin'], '| машина в имени ОПО', svod['opo'])
print('карточек с ИНН:', len(kartochki), '| предприятий:', len({k['inn'] for k in kartochki}))
vidy = {}
for k in kartochki: vidy[k['tip']] = vidy.get(k['tip'], 0) + 1
print('по типам:', vidy)
print('файл:', out)
