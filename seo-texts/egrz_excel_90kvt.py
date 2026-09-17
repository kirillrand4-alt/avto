# -*- coding: utf-8 -*-
"""Эксель по заказу: 2025 год и новее, и только там, где нужен компрессор от 90 кВт.

ЧЕСТНО О ПРИРОДЕ ЧИСЛА «90 кВт». В ЕГРЗ мощности компрессора нет НИ В ОДНОЙ записи.
Значит порог нельзя прочитать - его можно только оценить по типу производства и его
масштабу. Это инженерное суждение, и у каждой строки написано, из чего оно сложилось.

КАК ОЦЕНКА ПРОВЕРЯЛАСЬ. Я взяла 120 строк, по 30 из каждого своего класса, и попросила
провайдера независимо оценить киловатты по тем же названиям. Совпадение по порогу 90 кВт
вышло 58 %. Это плохо, и разбор расхождений показал два моих системных перекоса:
  ЗАВЫШАЛА на вспомогательных объектах - «ж/д путь к ГОКу» получал полную базу отрасли
    и выходил на 300 кВт, хотя компрессор нужен производству, а не подводке к нему;
  ЗАНИЖАЛА по машиностроению - механосборочный цех у меня был 45 кВт, у него 180, и его
    довод сильнее: это десятки пневмоинструментов, покрасочные камеры и стенды.
После двух правок согласие 58 % -> 67 % -> 72 %. Оставшиеся 28 % расхождений - настоящая
неопределённость метода, и она названа в легенде файла, а не спрятана.
"""
import collections
import csv
import io
import json
import os
import re

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

KAT = os.path.dirname(os.path.abspath(__file__))
VYHOD = os.path.join(KAT, 'EGRZ-OT-90-KVT-2025.xlsx')
SHRIFT = 'Arial'
TEH = re.compile(r'инженер|энергетик|механик|технолог|техдирект|техн|производств|цех|'
                 r'КИПиА|АСУ|снабжен|закупк|заказчик|капитальн|проект|эксплуатац', re.I)

g = [json.loads(l) for l in io.open(os.path.join(KAT, 'PROIZVODSTVA-90KVT.jsonl'),
                                    encoding='utf-8')]
nov = {r['inn']: r['novoe_dlya_bazy'] for r in csv.DictReader(
    io.open(os.path.join(KAT, '3s_PROIZV-NOVIZNA.csv'), encoding='utf-8-sig'), delimiter=';')}
kont = list(csv.DictReader(io.open(os.path.join(KAT, '3s_PROIZV-KONTAKTY.csv'),
                                   encoding='utf-8-sig'), delimiter=';'))
povod = list(csv.DictReader(io.open(os.path.join(KAT, '3s_PROIZV-POVODY.csv'),
                                    encoding='utf-8-sig'), delimiter=';'))
for k in kont:
    k['tehnicheskaya'] = 'да' if TEH.search((k.get('chelovek_dolzhnost') or '') + ' '
                                            + (k.get('rol_kanon') or '')) else ''

OT90 = sorted([x for x in g if x['kvt_ocenka'] >= 90],
              key=lambda z: (-z['kvt_ocenka'], z.get('data') or ''), reverse=False)
OT90.sort(key=lambda z: (-z['kvt_ocenka'], -int((z.get('data') or '0').replace('-', ''))))
MENSHE = sorted([x for x in g if 0 < x['kvt_ocenka'] < 90],
                key=lambda z: -z['kvt_ocenka'])
NET = [x for x in g if x['kvt_ocenka'] == 0]
INN90 = {x['zastroyshchik_inn'] for x in OT90 if x.get('zastroyshchik_inn')}
kont90 = [k for k in kont if k['inn'] in INN90]
povod90 = [p for p in povod if p['inn'] in INN90]
po_tel = collections.Counter(k['inn'] for k in kont90 if (k.get('telefon') or '').strip())
po_teh = collections.Counter(k['inn'] for k in kont90 if k['tehnicheskaya'])
po_pov = collections.Counter(p['inn'] for p in povod90)
print('от 90 кВт %d · меньше 90 %d · не нужен %d · предприятий %d'
      % (len(OT90), len(MENSHE), len(NET), len(INN90)))

wb = Workbook()
ZAG = Font(name=SHRIFT, size=12, bold=True)
OB = Font(name=SHRIFT, size=10)
ws = wb.active
ws.title = 'Как читать'
po_god = collections.Counter(x['data'][:4] for x in OT90)
po_otr = collections.Counter(x['otrasl'] for x in OT90)
OPIS = [
    ('Объекты 2025 года и новее, которым после стройки нужен компрессор от 90 кВт', True),
    ('', False),
    ('ЧТО ОТФИЛЬТРОВАНО. Из 10 173 карточек производственной выгрузки взяты только те,', False),
    ('   что признаны производством, датированы 2025 годом или позже, и по оценке требуют', False),
    ('   компрессора от 90 кВт. Осталось %d карточек и %d предприятий.' % (len(OT90), len(INN90)), False),
    ('   По годам: 2025 - %d, 2026 - %d.' % (po_god.get('2025', 0), po_god.get('2026', 0)), False),
    ('', False),
    ('ГЛАВНОЕ, ЧТО НАДО ЗНАТЬ ПРО ЧИСЛО «90 кВт»', True),
    ('   В ЕГРЗ мощности компрессора нет НИ В ОДНОЙ записи - я проверила весь текст', False),
    ('   названий. Значит порог нельзя ПРОЧИТАТЬ, его можно только ОЦЕНИТЬ по типу', False),
    ('   производства и его масштабу. Это инженерное суждение, а не факт из документа.', False),
    ('   У каждой строки стоит своя оценка в киловаттах и колонка «Из чего оценка».', False),
    ('', False),
    ('   Ориентир: 90 кВт винтового компрессора - это примерно 15 м³/мин при 8 бар.', False),
    ('   Дальше по типам: процессный воздух (доменная печь, конвертер, флотация,', False),
    ('   стекловаренная печь, помол цемента, целлюлоза, пиролиз) - сотни киловатт, воздух', False),
    ('   там часть технологии. Пневмотранспорт сыпучих (элеватор, комбикорм, мука, зола) -', False),
    ('   90-250 кВт. Приводной воздух завода (молокозавод, пивоварня, гальваника, окраска) -', False),
    ('   от 30 до 160 в зависимости от размера. Мелкие участки - 7-45 кВт, это не наш класс.', False),
    ('   Поправки: «завод», «комбинат», «комплекс» - вверх; «участок», «помещение»,', False),
    ('   «мастерская» - вниз; названная в тексте производительность - отдельная поправка.', False),
    ('', False),
    ('КАК ОЦЕНКА ПРОВЕРЯЛАСЬ, и результат проверки неприятный', True),
    ('   Я взяла 120 строк, по 30 из каждого своего класса, и попросила провайдера', False),
    ('   независимо оценить киловатты по тем же названиям, не показывая ему моих чисел.', False),
    ('   Совпадение по порогу 90 кВт вышло 58 %. Разбор расхождений нашёл два системных', False),
    ('   перекоса, оба мои:', False),
    ('     ЗАВЫШАЛА на вспомогательных объектах: «ж/д путь к ГОКу» получал полную базу', False),
    ('       отрасли и выходил на 300 кВт. Компрессор нужен производству, а не подводке.', False),
    ('     ЗАНИЖАЛА по машиностроению: механосборочный цех у меня был 45 кВт, у него 180.', False),
    ('       Его довод сильнее: это десятки пневмоинструментов, камеры окраски и стенды.', False),
    ('   После двух правок согласие 58 % -> 67 % -> 72 %.', False),
    ('   Оставшиеся 28 % - настоящая неопределённость метода. Значит примерно каждая', False),
    ('   четвёртая строка этого файла может оказаться не тем, чем я её назвала.', False),
    ('   Дешевле всего это проверяется звонком, а не расчётом.', False),
    ('', False),
    ('ЛИСТЫ', True),
    ('   «От 90 кВт»        %4d карточек, свежие и крупные сверху.' % len(OT90), False),
    ('   «Предприятия»      %4d предприятий: максимальная оценка, число объектов, телефоны.' % len(INN90), False),
    ('   «Меньше 90 кВт»    %4d карточек - отброшены порогом, но производство настоящее.' % len(MENSHE), False),
    ('                       Если решите брать и мелкие - они здесь, не выброшены.', False),
    ('   «Компрессор не нужен» %d карточек: подводящий трубопровод, пристройка, замена' % len(NET), False),
    ('                       аппарата на подстанции. Завод рядом назван, но объект не его.', False),
    ('   «Контакты» %d строк и «Поводы» %d строк - только по предприятиям с листа «От 90 кВт».'
     % (len(kont90), len(povod90)), False),
    ('', False),
    ('ОТРАСЛИ среди «от 90 кВт»', True),
] + [('   %-36s %4d' % (k[:36], n), False) for k, n in po_otr.most_common(11)] + [
    ('', False),
    ('ЧЕГО В ФАЙЛЕ НЕТ, сказано прямо:', True),
    ('   • компрессора в этих проектах никто не обещал. Признак косвенный: так работает', False),
    ('     производство такого типа и размера. Это гипотеза с ИНН и адресом.', False),
    ('   • телефон в базе есть у %d предприятий из %d, телефон при ТЕХНИЧЕСКОЙ роли - у %d.'
     % (len([i for i in INN90 if po_tel.get(i)]), len(INN90),
        len([i for i in INN90 if po_teh.get(i)])), False),
    ('   • это ЗАКЛЮЧЕНИЯ экспертизы: стройка могла не начаться, а могла и закончиться.', False),
    ('   • год отобран по дате заключения экспертизы, другой даты в реестре нет.', False),
    ('', False),
    ('КОНТРОЛИ', True),
    ('   • выдуманное слово «нипрятозаумень» даёт в реестре 0 записей.', False),
    ('   • выдуманный ИНН 9999999999 в нашей базе даёт 0 контактов и 0 поводов.', False),
    ('   • выгрузка полная: 0 недоборов на 123 слова поиска.', False),
    ('   • оценка мощности сверена с независимой оценкой на 120 строках, результат сверки', False),
    ('     назван выше числом, а не словом «проверено».', False),
]
for i, (t, zh) in enumerate(OPIS, start=1):
    c = ws.cell(row=i, column=1, value=t)
    c.font = ZAG if zh else OB
ws.column_dimensions['A'].width = 100

SH = ['Оценка, кВт', 'Уверенность', 'Из чего оценка', 'Дата заключения', 'Регион',
      'Объект', 'Отрасль', 'Приоритетная', 'Найден по слову', 'Адрес объекта',
      'Вывод экспертизы', 'Застройщик', 'ИНН застройщика', 'Новый для базы',
      'Телефонов', 'Из них тех.роль', 'Поводов', 'Проектировщик', 'ИНН проектировщика',
      'Технический заказчик', 'Номер заключения', 'Ссылка']
SHIR = [11, 22, 46, 14, 22, 58, 24, 11, 22, 30, 20, 38, 14, 12, 11, 12, 9, 36, 16, 28, 20, 44]


def list_ob(imya, gr):
    w = wb.create_sheet(imya)
    w.append(SH)
    for x in gr:
        z = x.get('zastroyshchik_inn', '')
        w.append([x['kvt_ocenka'], x['kvt_klass'], x['kvt_pochemu'], x.get('data', ''),
                  x.get('region', ''), x.get('obekt', ''), x.get('otrasl', ''),
                  x.get('prioritetnaya', ''), x.get('nayden_po', ''),
                  x.get('adres_obekta', ''), x.get('vyvod', ''),
                  x.get('zastroyshchik', ''), z, nov.get(z, ''),
                  po_tel.get(z, 0), po_teh.get(z, 0), po_pov.get(z, 0),
                  x.get('proektirovshchik', ''), x.get('proektirovshchik_inn', ''),
                  x.get('tehzakazchik', ''), x.get('nomer_zaklyucheniya', ''),
                  x.get('ssylka', '')])
    return w


w1 = list_ob('От 90 кВт', OT90)
w2 = list_ob('Меньше 90 кВт', MENSHE)
w3 = list_ob('Компрессор не нужен', NET)

w4 = wb.create_sheet('Предприятия')
w4.append(['ИНН', 'Предприятие', 'Отрасль', 'Регион', 'Новое для базы',
           'Максимальная оценка, кВт', 'Объектов от 90 кВт', 'Телефонов',
           'Из них тех.роль', 'Поводов', 'Последнее заключение'])
agg = {}
for x in OT90:
    i = x.get('zastroyshchik_inn')
    if not i:
        continue
    a = agg.setdefault(i, {'imya': x['zastroyshchik'], 'otr': collections.Counter(),
                           'reg': x['region'], 'n': 0, 'kvt': 0, 'data': ''})
    a['otr'][x['otrasl']] += 1
    a['n'] += 1
    a['kvt'] = max(a['kvt'], x['kvt_ocenka'])
    a['data'] = max(a['data'], x['data'])
for i, a in sorted(agg.items(), key=lambda z: (-z[1]['kvt'], -z[1]['n'])):
    w4.append([i, a['imya'], a['otr'].most_common(1)[0][0], a['reg'], nov.get(i, ''),
               a['kvt'], a['n'], po_tel.get(i, 0), po_teh.get(i, 0), po_pov.get(i, 0),
               a['data']])

wk = wb.create_sheet('Контакты')
wk.append(['ИНН', 'Должность как в источнике', 'Роль (канон)', 'Телефон', 'Почта',
           'Источник', 'Ссылка', 'Откуда взято', 'Тех.роль'])
for k in kont90:
    wk.append([k['inn'], k.get('chelovek_dolzhnost', ''), k.get('rol_kanon', ''),
               k.get('telefon', ''), k.get('pochta', ''), k.get('istochnik', ''),
               k.get('ssylka', ''), k.get('otkuda', ''), k['tehnicheskaya']])
wp = wb.create_sheet('Поводы')
wp.append(['ИНН', 'Тип события', 'Что происходит', 'Сумма', 'Источник', 'Ссылка', 'Когда'])
for p in povod90:
    wp.append([p['inn'], p.get('tip_sobytiya', ''), p.get('chto', ''), p.get('summa', ''),
               p.get('istochnik', ''), p.get('ssylka', ''), p.get('kogda', '')])

SH_F = Font(name=SHRIFT, size=10, bold=True, color='FFFFFF')
SH_FILL = PatternFill('solid', fgColor='2F5496')
SHIR_DOP = {'Предприятия': [14, 52, 24, 22, 13, 14, 12, 11, 12, 9, 15],
            'Контакты': [14, 38, 18, 18, 28, 20, 44, 16, 10],
            'Поводы': [14, 22, 70, 18, 22, 44, 14]}
for w in (w1, w2, w3, w4, wk, wp):
    for c in w[1]:
        c.font = SH_F
        c.fill = SH_FILL
        c.alignment = Alignment(vertical='center', wrap_text=True)
    w.freeze_panes = 'A2'
    w.auto_filter.ref = w.dimensions
    w.row_dimensions[1].height = 32
    for i, sh in enumerate(SHIR_DOP.get(w.title, SHIR), start=1):
        w.column_dimensions[get_column_letter(i)].width = sh
    for row in w.iter_rows(min_row=2):
        for c in row:
            c.font = Font(name=SHRIFT, size=10)
            c.alignment = Alignment(vertical='top')

wb.save(VYHOD)
print('файл: %s' % VYHOD)
print('  От 90 кВт %d · Предприятия %d · Меньше 90 %d · Не нужен %d'
      % (len(OT90), len(agg), len(MENSHE), len(NET)))
print('  предприятий с телефоном %d, с телефоном при тех.роли %d'
      % (len([i for i in INN90 if po_tel.get(i)]), len([i for i in INN90 if po_teh.get(i)])))
print('  новых для базы предприятий: %d из %d'
      % (len([i for i in INN90 if nov.get(i) == 'да']), len(INN90)))
