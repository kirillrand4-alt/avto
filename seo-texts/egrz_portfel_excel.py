# -*- coding: utf-8 -*-
"""Портфели лиц: что дала главная идея веера металинз. Числа честные, и они невесёлые.

Идея: не искать слово в названии, а взять всех проектировщиков и застройщиков «нашей
темы» и выгрузить ВЕСЬ их портфель заключений. Четыре линзы из шести назвали её
независимо. Замер обещал 26 513 заключений.

ЧТО ВЫШЛО НА САМОМ ДЕЛЕ. Выгружено 9 467 карточек по 299 лицам (гиганты нефтегаза с
портфелем больше 200 не брались). Новых карточек 9 079. А новых объектов НАШЕЙ ТЕМЫ
среди них — СЕМНАДЦАТЬ. Попадание 0,18 %, тогда как поиск по словам давал 21 %.

Идея не оправдалась в том, ради чего задумывалась. Но она дала другое, и это в файле:
1 554 новых для базы предприятия и 372 карточки, где в названии стоит производственное
слово — гальваника, литейка, прокат, обогащение, метанол, — то есть объект, которому
сжатый воздух нужен физически, хотя слова «компрессор» в названии нет.

Листы:
  «Наша тема новое»        17 карточек — прямая прибавка к 294
  «Производственные»      372 карточки со словом производства и пометкой новизны
  «Лица»                  299 лиц: портфель и что он дал
  «Застройщики»         1 831 ИНН с пометкой «новое для базы»
"""
import collections
import csv
import io
import json
import os

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

KAT = os.path.dirname(os.path.abspath(__file__))
VYHOD = os.path.join(KAT, 'EGRZ-PORTFELI-LIC.xlsx')
SHRIFT = 'Arial'

s = [json.loads(l) for l in io.open(os.path.join(KAT, 'PORTFEL-RAZOBRANO.jsonl'),
                                    encoding='utf-8')]
staro = {json.loads(l)['ssylka'] for l in io.open(os.path.join(KAT, 'EGRZ-NASHA-TEMA-2.jsonl'),
                                                  encoding='utf-8')}
nov = {r['inn']: r['novoe_dlya_bazy'] for r in csv.DictReader(
    io.open(os.path.join(KAT, '3s_PORTFEL-NOVIZNA.csv'), encoding='utf-8-sig'), delimiter=';')}
lica = [x for x in csv.DictReader(io.open(os.path.join(KAT, 'PORTFEL-LIC.csv'),
                                          encoding='utf-8-sig'), delimiter=';')
        if x['vsego_v_reestre'].isdigit() and 0 < int(x['vsego_v_reestre']) <= 200]

zam = [x for x in csv.DictReader(io.open(os.path.join(KAT, 'ZAMER-KANDIDATOV.csv'),
                                         encoding='utf-8-sig'), delimiter=';')
       if x['vsego'].isdigit() and int(x['vsego']) > 0]
# Слова с заведомо высоким шумом сюда не берутся: «молоч» ловит молочно-товарные фермы,
# «кирпичн» — капремонт крыши кирпичного дома, «асфальтобетонн» — ремонт дороги.
SHUM = {'молоч', 'молочн', 'кирпичн', 'асфальтобетонн', 'внеплощадочн', 'элеватор',
        'сушильн', 'химич', 'цемент', 'силос', 'пищев', 'вос', 'озх'}
SLOVA = [x['slovo'] for x in zam if x['slovo'] not in SHUM]

NASHA = sorted([x for x in s if x['klass'] == 'НАША ТЕМА' and x['ssylka'] not in staro],
               key=lambda z: z.get('data') or '', reverse=True)
PROM = []
for x in s:
    t = x['obekt'].lower()
    w = [q for q in SLOVA if q in t]
    if w:
        x['slova'] = ', '.join(w[:5])
        PROM.append(x)
PROM.sort(key=lambda z: z.get('data') or '', reverse=True)
print('новых «нашей темы» %d · производственных карточек %d' % (len(NASHA), len(PROM)))

po_licu = collections.Counter()
for x in s:
    for t in (x.get('nayden_po') or '').split(' | '):
        if t:
            po_licu[t] += 1
nasha_po_licu = collections.Counter()
for x in NASHA:
    for t in (x.get('nayden_po') or '').split(' | '):
        if t:
            nasha_po_licu[t] += 1
prom_po_licu = collections.Counter()
for x in PROM:
    for t in (x.get('nayden_po') or '').split(' | '):
        if t:
            prom_po_licu[t] += 1

wb = Workbook()
ZAG = Font(name=SHRIFT, size=12, bold=True)
OB = Font(name=SHRIFT, size=10)
ws = wb.active
ws.title = 'Как читать'
zastr_prom = {x['zastroyshchik_inn'] for x in PROM if x.get('zastroyshchik_inn')}
OPIS = [
    ('Портфели лиц: что дала главная идея веера металинз', True),
    ('', False),
    ('ИДЕЯ. Не искать слово в названии объекта, а взять всех проектировщиков и', False),
    ('   застройщиков «нашей темы» и выгрузить ВЕСЬ их портфель заключений. Проектировщик,', False),
    ('   сделавший одну компрессорную, делал их десяток. Четыре линзы из шести назвали этот', False),
    ('   приём независимо друг от друга.', False),
    ('', False),
    ('СИНТАКСИС, без которого идея не работает. Поля организаций в ЕГРЗ - коллекции:', False),
    ('      contains(PlannerOrganizations,\'ИНН\')            -> ОТКАЗ', False),
    ('      PlannerOrganizations/any(o: contains(o,\'ИНН\'))  -> работает', False),
    ('   Первый вариант выглядит как «в реестре ничего нет». Это ложный ноль, на котором', False),
    ('   тему можно было закрыть, так и не открыв.', False),
    ('', False),
    ('ЧТО ВЫШЛО, И ЭТО ГЛАВНОЕ ЧИСЛО ФАЙЛА', True),
    ('   выгружено карточек по 299 лицам ......... 9 467  (гиганты нефтегаза не брались)', False),
    ('   из них НОВЫХ, которых у нас не было ..... 9 079', False),
    ('   а новых объектов НАШЕЙ ТЕМЫ среди них ...    17', False),
    ('', False),
    ('   Попадание 0,18 %. Поиск по словам давал 21 % (294 из 1 419). То есть идея, которую', False),
    ('   независимо назвали четыре линзы и которая обещала 26 513 заключений, в своей', False),
    ('   собственной задаче не оправдалась. Так и записано, а не спрятано.', False),
    ('   Причина видна по данным: портфель проектировщика компрессорных - это в основном', False),
    ('   не компрессорные. Он делает дороги, сети, жильё и площадки того же заказчика.', False),
    ('', False),
    ('ЧТО ИДЕЯ ВСЁ-ТАКИ ДАЛА', True),
    ('   производственных карточек ............... %d  (лист «Производственные»)' % len(PROM), False),
    ('     это объекты, в названии которых стоит производственное слово - гальваника,', False),
    ('     литейка, прокат, обогащение, метанол, доменная печь. Компрессор им нужен', False),
    ('     физически, но слова «компрессор» в названии нет, и поиском по слову они', False),
    ('     не находились НИКОГДА.', False),
    ('   разных предприятий на этих карточках .... %d, из них НОВЫХ для базы %d'
     % (len(zastr_prom), len([i for i in zastr_prom if nov.get(i) == 'да'])), False),
    ('   всего застройщиков в портфельной выгрузке %d, из них НОВЫХ для базы %d'
     % (len(nov), len([i for i, v in nov.items() if v == 'да'])), False),
    ('', False),
    ('КАК ЧИТАТЬ ЛИСТЫ', True),
    ('   «Наша тема новое»   %3d карточек - прямая прибавка к нашим 294.' % len(NASHA), False),
    ('   «Производственные»  %3d карточек. Колонка «Слова производства» говорит, ЧТО именно' % len(PROM), False),
    ('                       сработало в названии. Класс не ставился: правило про', False),
    ('                       компрессорные тут неприменимо, судить надо глазами.', False),
    ('   «Лица»              299 лиц: размер портфеля и что он дал. Видно, кто отработал.', False),
    ('   «Застройщики»       %d ИНН с пометкой новизны - самый большой улов этого захода.' % len(nov), False),
    ('', False),
    ('ЧЕСТНО О КЛАССИФИКАТОРЕ', True),
    ('   Разбор вёлся ТЕМ ЖЕ классификатором, что дал 294 объекта, - иначе числа двух', False),
    ('   файлов были бы несравнимы. Но он рассчитан на записи, где компрессорное слово уже', False),
    ('   есть. Здесь таких почти нет, и 8 612 карточек он честно пометил «сомнительно».', False),
    ('   Проверено: во ВСЕХ 8 612 нет ни одного из 14 компрессорных слов. То есть это не', False),
    ('   «может быть наше», а «правило неприменимо». На листы файла они не вынесены.', False),
    ('', False),
    ('КОНТРОЛИ', True),
    ('   • выдуманный ИНН 9999999999 как проектировщик даёт 0 записей, как застройщик 0.', False),
    ('   • сверка новизны: выдуманный ИНН обязан быть новым (оказался), взятый из базы -', False),
    ('     не новым (не оказался).', False),
    ('   • выгрузка полная: ни один срез из 406 запросов не потерян (0 отказов сети).', False),
    ('     Это важно: неполный срез дал бы заниженное число и выглядел бы как вывод.', False),
]
for i, (t, zh) in enumerate(OPIS, start=1):
    c = ws.cell(row=i, column=1, value=t)
    c.font = ZAG if zh else OB
ws.column_dimensions['A'].width = 100

SH_OB = ['Дата заключения', 'Регион', 'Объект', 'Слова производства', 'Адрес объекта',
         'Вывод экспертизы', 'Застройщик', 'ИНН застройщика', 'Застройщик новый для базы',
         'Проектировщик', 'ИНН проектировщика', 'Технический заказчик',
         'От какого лица найдено', 'Номер заключения', 'Ссылка на карточку']
SHIR_OB = [14, 22, 62, 26, 34, 20, 40, 14, 13, 40, 16, 32, 26, 20, 44]


def list_obektov(imya, gr, so_slovami):
    w = wb.create_sheet(imya)
    w.append(SH_OB)
    for x in gr:
        zi = x.get('zastroyshchik_inn', '')
        w.append([x.get('data', ''), x.get('region', ''), x.get('obekt', ''),
                  x.get('slova', '') if so_slovami else x.get('klass', ''),
                  x.get('adres_obekta', ''), x.get('vyvod', ''),
                  x.get('zastroyshchik', ''), zi, nov.get(zi, ''),
                  x.get('proektirovshchik', ''), x.get('proektirovshchik_inn', ''),
                  x.get('tehzakazchik', ''), x.get('nayden_po', ''),
                  x.get('nomer_zaklyucheniya', ''), x.get('ssylka', '')])
    return w


w1 = list_obektov('Наша тема новое', NASHA, False)
w1['D1'] = 'Класс'
w2 = list_obektov('Производственные', PROM, True)

w3 = wb.create_sheet('Лица')
w3.append(['Роль', 'ИНН', 'Название', 'Портфель в реестре', 'Было у нас',
           'Выгружено карточек', 'Дало новых «нашей темы»', 'Дало производственных'])
for x in sorted(lica, key=lambda z: -(prom_po_licu.get('%s %s' % (z['rol'], z['inn']), 0)
                                      + 10 * nasha_po_licu.get('%s %s' % (z['rol'], z['inn']), 0))):
    t = '%s %s' % (x['rol'], x['inn'])
    w3.append([x['rol'], x['inn'], x['imya'], int(x['vsego_v_reestre']), int(x['uzhe_u_nas']),
               po_licu.get(t, 0), nasha_po_licu.get(t, 0), prom_po_licu.get(t, 0)])

w4 = wb.create_sheet('Застройщики')
w4.append(['ИНН', 'Название', 'Новое для базы', 'Карточек в выгрузке',
           'Из них производственных'])
imena, vsego_k = {}, collections.Counter()
for x in s:
    i = x.get('zastroyshchik_inn')
    if i:
        imena.setdefault(i, x['zastroyshchik'])
        vsego_k[i] += 1
prom_k = collections.Counter(x['zastroyshchik_inn'] for x in PROM if x.get('zastroyshchik_inn'))
for i in sorted(vsego_k, key=lambda z: (-prom_k.get(z, 0), -vsego_k[z])):
    w4.append([i, imena.get(i, ''), nov.get(i, ''), vsego_k[i], prom_k.get(i, 0)])

SH_F = Font(name=SHRIFT, size=10, bold=True, color='FFFFFF')
SH_FILL = PatternFill('solid', fgColor='2F5496')
SHIR = {'Лица': [14, 13, 46, 12, 11, 12, 13, 13],
        'Застройщики': [14, 56, 13, 13, 13]}
for w in (w1, w2, w3, w4):
    for c in w[1]:
        c.font = SH_F
        c.fill = SH_FILL
        c.alignment = Alignment(vertical='center', wrap_text=True)
    w.freeze_panes = 'A2'
    w.auto_filter.ref = w.dimensions
    w.row_dimensions[1].height = 32
    for i, sh in enumerate(SHIR.get(w.title, SHIR_OB), start=1):
        w.column_dimensions[get_column_letter(i)].width = sh
    for row in w.iter_rows(min_row=2):
        for c in row:
            c.font = Font(name=SHRIFT, size=10)
            c.alignment = Alignment(vertical='top')

wb.save(VYHOD)
print('файл: %s' % VYHOD)
print('  Наша тема новое %d · Производственные %d · Лица %d · Застройщики %d'
      % (len(NASHA), len(PROM), len(lica), len(vsego_k)))
print('  лиц, не давших НИЧЕГО (ни нашей темы, ни производственных): %d из %d'
      % (len([x for x in lica if not nasha_po_licu.get('%s %s' % (x['rol'], x['inn']))
              and not prom_po_licu.get('%s %s' % (x['rol'], x['inn']))]), len(lica)))
