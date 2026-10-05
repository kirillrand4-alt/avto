# -*- coding: utf-8 -*-
"""Объединение моего среза со срезом 2-й сессии. Подтверждённое дважды видно отдельно.

Правило владельца: источники НАКАПЛИВАЮТСЯ, а не заменяются, и подтверждённое двумя
разборами должно быть отличимо от подтверждённого одним. Две сессии собрали базу под
Мейер параллельно и независимо - это и есть две независимые проверки одного материала.

ЧТО ПОКАЗАЛА СВЕРКА ДО СЛИЯНИЯ:
  у меня 265 пар «ИНН + номер» по 182 компаниям
  у 2-й сессии 144 пары по 105 компаниям
  общих 122 - вот они и подтверждены дважды
  только у меня 143, только у неё 22

Её уникальные - в основном роль «продажи», которую я отсекала по списку ЛПР. Но там же
нашлась и настоящая потеря с моей стороны: «Кованов Александр, главный инженер» с мобильным
из tender.pro. Поэтому слияние, а не выбор одного файла.
"""
import collections
import os
import re

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

KAT = os.path.dirname(os.path.abspath(__file__))
VYHOD = os.path.join(KAT, 'MEYER-LPR-DOZVON-OBE-SESSII.xlsx')
SHRIFT = 'Arial'


def cifry(s):
    c = re.sub(r'\D', '', str(s or ''))
    if len(c) == 11 and c.startswith('8'):
        c = '7' + c[1:]
    elif len(c) == 10 and c.startswith('9'):
        c = '7' + c
    return c


def krasivo(c):
    return ('+%s (%s) %s-%s-%s' % (c[0], c[1:4], c[4:7], c[7:9], c[9:11])
            if len(c) == 11 else c)


zapisi = {}


def dobavit(klyuch, d, sessiya):
    z = zapisi.get(klyuch)
    if not z:
        d['sessii'] = [sessiya]
        zapisi[klyuch] = d
        return
    z['sessii'].append(sessiya)
    # Накапливаем, а не заменяем: пустое поле у одной сессии заполняется из другой,
    # ссылки складываются.
    for p in ('fio', 'dolzhnost', 'rol', 'nazvanie', 'region', 'sayt', 'okved',
              'vyruchka', 'segment'):
        if not str(z.get(p) or '').strip() and str(d.get(p) or '').strip():
            z[p] = d[p]
    for p in ('istochnik', 'ssylka'):
        a, b = str(z.get(p) or ''), str(d.get(p) or '')
        if b and b not in a:
            z[p] = (a + ' | ' + b).strip(' |')


# ---------- мой файл
w1 = load_workbook(os.path.join(KAT, 'MEYER-LPR-MOBILNYE-I-DOBAVOCHNYE.xlsx'))
for lst in ('Мобильные и добавочные', 'Элеваторы запасной ход'):
    r = list(w1[lst].values)
    sh = [str(x) for x in r[0]]
    for row in r[1:]:
        x = dict(zip(sh, row))
        c = cifry(x.get('Телефон'))
        if len(c) != 11 and 'добавочн' not in str(x.get('Тип номера') or ''):
            continue
        dobavit((str(x.get('ИНН')), c or str(x.get('Телефон'))), {
            'inn': str(x.get('ИНН') or ''), 'nomer': krasivo(c) or str(x.get('Телефон') or ''),
            'tip': x.get('Тип номера') or '', 'fio': x.get('ФИО') or '',
            'dolzhnost': '', 'rol': x.get('Класс роли') or '',
            'rol_istochnik': x.get('Роль как в источнике') or '',
            'nazvanie': x.get('Название') or '', 'segment': x.get('Сегменты') or '',
            'region': '', 'sayt': '', 'okved': '', 'vyruchka': '',
            'istochnik': x.get('Источник') or '', 'ssylka': x.get('Ссылка на источник') or '',
            'list': lst,
        }, '3-я сессия')

# ---------- файл 2-й сессии
w2 = load_workbook(os.path.join(KAT, 'meyer-lpr-tel-0510.xlsx'))
r = list(w2['Контакты'].values)
sh = [str(x) for x in r[0]]
for row in r[1:]:
    x = dict(zip(sh, row))
    for pole, tip in (('Мобильный', 'мобильный'),
                      ('Рабочий с добавочным', 'городской с добавочным')):
        v = str(x.get(pole) or '').strip()
        if not v:
            continue
        c = cifry(v)
        dobavit((str(x.get('ИНН')), c or v), {
            'inn': str(x.get('ИНН') or ''), 'nomer': krasivo(c) or v, 'tip': tip,
            'fio': x.get('ФИО') or '', 'dolzhnost': x.get('Должность') or '',
            'rol': x.get('Роль') or '', 'rol_istochnik': x.get('Должность') or '',
            'nazvanie': x.get('Название') or '', 'segment': x.get('Сегмент') or '',
            'region': x.get('Регион') or '', 'sayt': x.get('Сайт') or '',
            'okved': x.get('Основной ОКВЭД') or '', 'vyruchka': x.get('Выручка, руб') or '',
            'istochnik': x.get('Источник') or '', 'ssylka': x.get('Ссылка на источник') or '',
            'list': 'Контакты 2С',
        }, '2-я сессия')

VSE = list(zapisi.values())
for z in VSE:
    z['sessii'] = sorted(set(z['sessii']))
    z['sessiy'] = len(z['sessii'])
    z['kem'] = ' + '.join(z['sessii'])
dvazhdy = [z for z in VSE if z['sessiy'] >= 2]
print('строк всего: %d' % len(VSE))
print('  подтверждено ДВУМЯ сессиями: %d' % len(dvazhdy))
print('  только 3-я сессия: %d' % len([z for z in VSE if z['kem'] == '3-я сессия']))
print('  только 2-я сессия: %d' % len([z for z in VSE if z['kem'] == '2-я сессия']))
print('  компаний: %d' % len({z['inn'] for z in VSE}))
print('  с именем человека: %d' % len([z for z in VSE if str(z['fio'] or '').strip()
                                       and str(z['fio']) != 'None']))

VSE.sort(key=lambda z: (-z['sessiy'], z['rol'] or 'я', z['nazvanie'] or ''))

wb = Workbook()
ZAG = Font(name=SHRIFT, size=12, bold=True)
OB = Font(name=SHRIFT, size=10)
ws = wb.active
ws.title = 'Как читать'
po_roli = collections.Counter(str(z['rol'] or 'не указана')[:34] for z in VSE)
OPIS = [
    ('Дозваниваемые ЛПР: объединение двух сессий, собиравших базу параллельно', True),
    ('', False),
    ('ЗАЧЕМ ЭТОТ ФАЙЛ. Две сессии собрали базу под Мейер независимо друг от друга и', False),
    ('   положили результаты на дроп в один час. Это не дубль работы, а две независимые', False),
    ('   проверки одного материала - значит из них можно взять то, чего нет поодиночке.', False),
    ('', False),
    ('ЧТО ПОКАЗАЛА СВЕРКА ДО СЛИЯНИЯ', True),
    ('   у 3-й сессии ...... 265 пар «ИНН + номер» по 182 компаниям', False),
    ('   у 2-й сессии ...... 144 пары по 105 компаниям', False),
    ('   общих ............. 122', False),
    ('   только у 3-й ...... 143', False),
    ('   только у 2-й ......  22', False),
    ('', False),
    ('ЧИСЛА ОБЪЕДИНЁННОГО ФАЙЛА', True),
    ('   строк ............................... %d' % len(VSE), False),
    ('   подтверждено ДВУМЯ сессиями ......... %d' % len(dvazhdy), False),
    ('   только 3-я сессия ................... %d'
     % len([z for z in VSE if z['kem'] == '3-я сессия']), False),
    ('   только 2-я сессия ................... %d'
     % len([z for z in VSE if z['kem'] == '2-я сессия']), False),
    ('   разных компаний ..................... %d' % len({z['inn'] for z in VSE}), False),
    ('   с именем человека ................... %d'
     % len([z for z in VSE if str(z['fio'] or '').strip() and str(z['fio']) != 'None']), False),
    ('', False),
    ('   Колонка «Сессий» и есть провенанс: 2 значит номер нашли двумя разными ходами.', False),
    ('   Подтверждённое дважды стоит сверху и должно звониться первым.', False),
    ('', False),
    ('ЧТО НАШЛОСЬ У СОСЕДКИ И ЧЕГО НЕ БЫЛО У МЕНЯ', True),
    ('   Её 22 уникальные строки - в основном роль «продажи», которую я отсекала по вашему', False),
    ('   списку ЛПР. Но там же нашлась и моя потеря: «Кованов Александр, главный инженер»', False),
    ('   с мобильным, источник tender.pro. Он в файле есть.', False),
    ('   Причина потери: её сбор шёл по сайтам компаний и подписям со страниц, мой - по', False),
    ('   нашей базе phone_contacts. Это разные ходы, и они дают разный улов.', False),
    ('', False),
    ('ПО РОЛЯМ', True),
] + [('   %-36s %5d' % (k, n), False) for k, n in po_roli.most_common(12)] + [
    ('', False),
    ('ЧТО С ЭТИМ ДЕЛАТЬ', True),
    ('   Звонить сверху вниз: сперва подтверждённые двумя сессиями, потом остальные.', False),
    ('   Оба исходных файла лежат на дропе и не тронуты: MEYER-LPR-MOBILNYE-I-DOBAVOCHNYE.xlsx', False),
    ('   и meyer-lpr-tel-0510.xlsx.', False),
]
for i, (t, zh) in enumerate(OPIS, start=1):
    c = ws.cell(row=i, column=1, value=t)
    c.font = ZAG if zh else OB
ws.column_dimensions['A'].width = 100

KOL = [('Сессий', 'sessiy', 8), ('Кем найден', 'kem', 26), ('Тип номера', 'tip', 22),
       ('Телефон', 'nomer', 22), ('ФИО', 'fio', 28), ('Должность', 'dolzhnost', 30),
       ('Класс роли', 'rol', 30), ('Роль как в источнике', 'rol_istochnik', 36),
       ('Название', 'nazvanie', 42), ('ИНН', 'inn', 13), ('Регион', 'region', 22),
       ('Сайт', 'sayt', 26), ('Основной ОКВЭД', 'okved', 16),
       ('Выручка, руб', 'vyruchka', 15), ('Сегмент', 'segment', 44),
       ('Источник', 'istochnik', 28), ('Ссылка на источник', 'ssylka', 58)]
w = wb.create_sheet('Дозваниваемые ЛПР')
w.append([a for a, _, _ in KOL])
for z in VSE:
    r = []
    for _, k, _s in KOL:
        v = z.get(k, '')
        if v is None or str(v) == 'None':
            v = ''
        r.append(v)
    w.append(r)

SH_F = Font(name=SHRIFT, size=10, bold=True, color='FFFFFF')
SH_FILL = PatternFill('solid', fgColor='2F5496')
for c in w[1]:
    c.font = SH_F
    c.fill = SH_FILL
    c.alignment = Alignment(vertical='center', wrap_text=True)
w.freeze_panes = 'A2'
w.auto_filter.ref = w.dimensions
w.row_dimensions[1].height = 32
for i, (_, _, s) in enumerate(KOL, start=1):
    w.column_dimensions[get_column_letter(i)].width = s
ZELEN = PatternFill('solid', fgColor='E2EFDA')
for row in w.iter_rows(min_row=2):
    for c in row:
        c.font = Font(name=SHRIFT, size=10)
        c.alignment = Alignment(vertical='top')
    if row[0].value == 2:
        for c in row:
            c.fill = ZELEN

wb.save(VYHOD)
print('\nфайл: %s' % VYHOD)
print('  подтверждённые двумя сессиями подсвечены зелёным и стоят сверху')
