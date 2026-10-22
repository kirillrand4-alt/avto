import openpyxl, json, re, collections
D="/root/.claude/uploads/eb90afa2-9dfc-57d7-acb9-abafd7969908/"; F=D+'1059e4ff-FSA-MEYER-v4-Meyer.xlsx'
S="/tmp/claude-0/-home-user-avto/eb90afa2-9dfc-57d7-acb9-abafd7969908/scratchpad/"
inns=json.load(open(S+'fsa/inns.json'))['meyer']; nov=set(json.load(open(S+'fsa/fsa-v4-novye.json'))['meyer'])
import sys; Y=sys.argv[1] if len(sys.argv)>1 else 'A'; A={i for i in nov if inns[i][0]==Y}
wb=openpyxl.load_workbook(F,read_only=True)
ws=wb['Досье']; it=ws.iter_rows(values_only=True); h=list(next(it)); ix={k:i for i,k in enumerate(h)}
g=lambda r,k:(str(r[ix[k]]).strip() if r[ix[k]] not in (None,'','—') else '')
EM=re.compile(r"[\w.+-]+@[\w-]+(?:\.[\w-]+)+")
comp=[]; c=collections.Counter(); keep=set()
for r in it:
    inn=g(r,'ИНН')
    if inn not in A: continue
    ok=g(r,'ОКВЭД основной')
    if ok[:2] not in ('10','11','01','03'):
        c['вон: не наш профиль']+=1; continue
    posr={e.lower() for e in EM.findall(g(r,'Посредник? (оформитель деклараций)'))}
    best=g(r,'Конвейер: Письмо: e-mail').lower()
    fio=g(r,'Конвейер: Письмо: кому (ФИО)'); dolzh=g(r,'Конвейер: Письмо: должность') or g(r,'Конвейер: Письмо: ступень лесенки')
    if not best:
        dek=[e.lower() for e in EM.findall(g(r,'Почты (из деклараций)')) if e.lower() not in posr]
        best=dek[0] if dek else ''; fio=''; dolzh='почта из декларации ФСА'
    if best in posr: best=''
    if not best:
        c['вон: только почта оформителя']+=1; continue
    pas={k[8:].strip().lower():g(r,k) for k in h if k.startswith('Паспорт:') and k not in ('Паспорт: уверенность','Паспорт: источники','Паспорт: откуда') and g(r,k)}
    pas['продукция по декларациям ФСА']=g(r,'Продукция (из деклараций по нашим основам)')[:600]
    pas['декларации']=('последняя %s; за 12 мес.: %s; действуют: %s'%(g(r,'Последняя декларация'),g(r,'За 12 мес.'),g(r,'Действуют сейчас')))
    for k,kk in (('Конвейер: Описание (по сайту)','описание по сайту'),('Цербер: виды предприятий','Цербер'),('Строят (ЕГРЗ 2025+)','строят (ЕГРЗ)'),('Поводы (новости)','повод')):
        if g(r,k): pas[kk]=g(r,k)[:400]
    site=g(r,'Конвейер: Сайт') or (g(r,'Ночь: сайт (ИНН на сайте)').split(' ')[0] if g(r,'Ночь: сайт (ИНН на сайте)') and '⚠' not in g(r,'Ночь: сайт (ИНН на сайте)') else '')
    comp.append({"inn":inn,"name":g(r,'Изготовитель'),"okved":ok,"region":g(r,'Регион'),"rev":float(g(r,'Выручка, руб') or 0) or None,
                 "site":site,"pochemu":g(r,'Почему'),"pasport":pas,"best":best,"best_fio":fio,"best_dolzh":dolzh,
                 "proval_meyer":g(r,'Битрикс Meyer: провал 2025–26')})
    keep.add(inn); c['в заливку']+=1
ws=wb['Контакты с сайта']; it=ws.iter_rows(values_only=True); h=list(next(it)); ix={k:i for i,k in enumerate(h)}
cont=[]
for r in it:
    inn=g(r,'ИНН / УНП'); em=g(r,'E-mail').lower()
    if inn not in keep or '@' not in em: continue
    lpr=g(r,'ЛПР по оборудованию'); role=g(r,'Роль')
    step=1 if lpr=='да' else 2 if lpr=='возможно' else 4 if role=='Продажи' else 3
    cont.append({"inn":inn,"email":em,"fio":g(r,'ФИО'),"dolzh":g(r,'Должность (как на сайте)'),"role":role,"step":step})
json.dump({"comp":comp,"cont":cont},open(S+('volna2-zaliv.json' if Y=='A' else 'volna2-zaliv-'+Y.lower()+'.json'),'w'),ensure_ascii=False)
print(dict(c),'контактов',len(cont))
