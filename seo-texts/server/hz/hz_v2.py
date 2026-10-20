# -*- coding: utf-8 -*-
"""Правки владельца поверх сборки hz_sborka (20.10): лишние листы убрать (как в его файле: Сводка, ОКВЭД не найден,
Не из 8 отраслей, Ликвидированы); во всех листах — только 8 отраслей; «Факты» — только закупки компрессоров и запчастей
к ним (без автомобильных/холодильных, без вакансий; «оборудование» — только если это закупка и не дубль закупки).

    python3 hz_v2.py <вход.xlsx> <выход.xlsx>
"""
import copy
import re
import sys

import openpyxl

ОТРАСЛИ = {'Пищевая промышленность и напитки', 'Машиностроение, металлообработка, приборостроение',
           'Фармацевтическая и медицинская промышленность', 'Химическая и нефтехимическая промышленность',
           'Строительная индустрия и стройматериалы', 'Деревообработка и мебель', 'Электроника и электротехника',
           'Лёгкая и текстильная промышленность'}
УБРАТЬ = ('Сводка', 'ОКВЭД не найден', 'Не из 8 отраслей', 'Ликвидированы')
ДА = re.compile(r'компрессор|винтов\w* блок|маслоотдел|сепаратор\w* (?:масл|воздуш)|ресивер|осушител\w* (?:воздух|сжат|рефриж|адсорб)|'
                r'сжат\w* воздух|\bДЭН-\d|\bЗИФ-?\d|\bПКС-?\d|\bВК-?\d{2}|\bНВ-?10|воздухоразделит', re.I)
НЕТ = re.compile(r'автомоб|камаз|\bмаз\b|тормоз|кондиционер|холодильн|сплит|фреон|хладон|вакансия|hh\.ru|машинист', re.I)
ЗАКУПКА = re.compile(r'ЭТП|ЕИС|tender|закупк|поставка|Росэлторг|РТС|ГПБ|zakupki', re.I)


def годится(тип, описание):
    о = str(описание or '')
    if тип not in ('закупка', 'план закупки', 'оборудование'):
        return False
    if тип == 'оборудование' and not ЗАКУПКА.search(о):
        return False
    return bool(ДА.search(о)) and not НЕТ.search(о)


def main(вход, выход):
    wb = openpyxl.load_workbook(вход)
    for n in УБРАТЬ:
        if n in wb.sheetnames:
            del wb[n]
    for ws in wb.worksheets:
        шапка = [c.value for c in ws[1]]
        if 'Отрасль' not in шапка:
            continue
        I = {k: i for i, k in enumerate(шапка)}
        keep = []
        for row in ws.iter_rows(min_row=2):
            v = [c.value for c in row]
            if not any(x not in (None, '') for x in v) or v[I['Отрасль']] not in ОТРАСЛИ:
                continue
            if ws.title == 'Факты' and not годится(v[I['Тип']], v[I['Описание']]):
                continue
            keep.append([(c.value, copy.copy(c._style)) for c in row])
        if ws.title == 'Факты':
            def текст(r):
                м = re.search(r'«(.{15,100}?)', str(r[I['Описание']][0] or ''))
                return re.sub(r'\W+', '', м.group(1)).lower() if м else None
            есть = {(r[I['ИНН']][0], текст(r)) for r in keep if r[I['Тип']][0] != 'оборудование' and текст(r)}
            keep = [r for r in keep if not (r[I['Тип']][0] == 'оборудование' and (r[I['ИНН']][0], текст(r)) in есть)]
            видел, k2 = set(), []
            for r in keep:
                k = (r[I['ИНН']][0], r[I['Номер']][0] or '', r[I['Дата']][0] or '', текст(r))
                if k not in видел:
                    видел.add(k)
                    k2.append(r)
            keep = k2
        было = ws.max_row - 1
        ws.delete_rows(2, ws.max_row)
        for n, r in enumerate(keep, 2):
            for k, (val, st) in enumerate(r, 1):
                c = ws.cell(row=n, column=k, value=val)
                c._style = st
        if ws.auto_filter.ref:
            ws.auto_filter.ref = 'A1:%s%d' % (openpyxl.utils.get_column_letter(ws.max_column), max(2, len(keep) + 1))
        print(ws.title, было, '->', len(keep))
    wb.save(выход)


if __name__ == '__main__':
    main(*sys.argv[1:3])
