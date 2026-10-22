# -*- coding: utf-8 -*-
# 22.10, владелец: «скинь актуальный список, поделив на 2 части — КЦ и Meyer, чтобы каждый весил не больше 30 МБ».
# Делит готовый FSA-MEYER-VSE (fsa_vse_xlsx.py) по колонке «Направление» листа «Работа»: строки «Работа»/«Досье» и
# «Контакты с сайта» (по ИНН) — в свой файл, цвет строки сохраняется; «Сводка» — в оба; листы Цербера — в Meyer.
# Пишет xlsxwriter-ом с общей таблицей строк: openpyxl write_only пишет каждую строку inline — Meyer-часть выходила 35 МБ
# («Досье» 23 МБ: повторы «Действующая компания», ОКВЭД, сегменты в каждой ячейке).
# Запуск: python3 fsa_razdelit.py <FSA-MEYER-VSE.xlsx> <выход-Meyer.xlsx> <выход-КЦ.xlsx>
import os, sys
import xlsxwriter
from openpyxl import load_workbook

ШИР = {'ИНН': 13, 'ИНН / УНП': 13, 'Компания': 40, 'Изготовитель': 40, 'Название': 40, 'Регион': 22, 'Сайт': 28,
       'Продукция (кратко)': 50, 'Паспорт (кратко)': 40, 'Пометки': 30, 'Битрикс': 30}


def main(вход, вых_м, вых_к):
    wb = load_workbook(вход, read_only=True)
    it = wb['Работа'].iter_rows(values_only=True)
    h = next(it)
    iН, iИ = h.index('Направление'), h.index('ИНН')
    кому = {str(r[iИ]): ('Meyer' if r[iН] == 'Meyer' else 'КЦ') for r in it}
    выход = {н: xlsxwriter.Workbook(п, {'strings_to_urls': False, 'strings_to_numbers': False, 'strings_to_formulas': False})
             for н, п in (('Meyer', вых_м), ('КЦ', вых_к))}
    шапка_ф = {н: в.add_format({'bold': True, 'bg_color': '#DDEBF7', 'text_wrap': True, 'valign': 'top'}) for н, в in выход.items()}
    жирн = {н: в.add_format({'bold': True}) for н, в in выход.items()}
    цвета = {н: {} for н in выход}

    def формат(н, rgb):
        if rgb not in цвета[н]:
            цвета[н][rgb] = выход[н].add_format({'bg_color': '#' + rgb[-6:]})
        return цвета[н][rgb]
    for лист in wb.sheetnames:
        строки = wb[лист].iter_rows()
        шапка = [c.value for c in next(строки)]
        ик = next((шапка.index(k) for k in ('ИНН', 'ИНН / УНП') if k in шапка), None)
        цели = ['Meyer', 'КЦ'] if лист == 'Сводка' or ик is None else ['Meyer'] if лист.startswith('Цербер') else None
        листы, ряд = {}, {}
        for н in (цели or ['Meyer', 'КЦ']):
            w = выход[н].add_worksheet(лист[:31])
            if лист == 'Сводка':
                w.set_column(0, 0, 110)
                w.set_column(1, 1, 16)
                w.write_row(0, 0, шапка, жирн[н])
            else:
                for i, k in enumerate(шапка):
                    w.set_column(i, i, ШИР.get(k, 18))
                w.write_row(0, 0, шапка, шапка_ф[н])
                w.freeze_panes(1, 2)
            листы[н], ряд[н] = w, 1
        for r in строки:
            if not r:
                for н in (цели or []):
                    ряд[н] += 1
                continue
            if цели:
                куда = цели
            else:
                н = кому.get(str(getattr(r[ик], 'value', '') or '')) if len(r) > ик else None
                if not н:
                    continue
                куда = [н]
            ф = getattr(r[0], 'fill', None)
            rgb = ф.fgColor.rgb if ф is not None and ф.fill_type == 'solid' and isinstance(ф.fgColor.rgb, str) else None
            знач = [getattr(c, 'value', None) for c in r]
            for н in куда:
                fmt = формат(н, rgb) if rgb else None
                for j, v in enumerate(знач):
                    if v is not None:
                        листы[н].write(ряд[н], j, v, fmt)
                    elif fmt:
                        листы[н].write_blank(ряд[н], j, None, fmt)
                ряд[н] += 1
        for н, w in листы.items():
            if лист != 'Сводка':
                w.autofilter(0, 0, max(ряд[н] - 1, 1), max(len(шапка) - 1, 0))
    for в in выход.values():
        в.close()
    for п in (вых_м, вых_к):
        print(os.path.basename(п), '%.1f МБ' % (os.path.getsize(п) / 2 ** 20))


if __name__ == '__main__':
    main(*sys.argv[1:4])
