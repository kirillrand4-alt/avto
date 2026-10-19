# -*- coding: utf-8 -*-
"""Отбор для рассылки из полного Excel КЦ (19.10, владелец: «по КЦ всё, что собрал, у кого есть почта, выручка любая
(но подписана, где известна), паспорт компании и т.д.»).

Берёт лист «Компании» полного файла (pilot_xlsx.py, KC_CEL=kc), оставляет компании с почтой, сохраняет все колонки
(выручка + «Откуда выручка», паспорт, ЛПР, адрес письма), сортирует: сначала с известной выручкой по убыванию, затем
без выручки. Лист «Контакты» — только по отобранным компаниям. Сводка — сколько и по каким срезам.

    python3 kc_otbor_xlsx.py <KC-1.xlsx> <KC-1-pochta.xlsx>
"""
import collections
import sys

import openpyxl
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter


def число(v):
    try:
        return float(v or 0)
    except (TypeError, ValueError):
        return 0.0


def main(вход, выход):
    wb = openpyxl.load_workbook(вход, read_only=True)
    it = wb['Компании'].iter_rows(values_only=True)
    шапка = list(next(it))
    к = {name: n for n, name in enumerate(шапка)}
    строки = [list(r) for r in it if r and any(v not in (None, '') for v in r)]
    всего = len(строки)
    паспорт = [n for n, name in enumerate(шапка) if str(name).startswith('Паспорт')]

    def с_почтой(r):
        return число(r[к['Почт']]) > 0 or bool(r[к['Письмо: e-mail']])

    отбор = [r for r in строки if с_почтой(r)]
    отбор.sort(key=lambda r: (число(r[к['Выручка, руб']]) <= 0, -число(r[к['Выручка, руб']])))
    инн = {str(r[к['ИНН / УНП']]) for r in отбор}

    сч = collections.Counter()
    for r in отбор:
        в = число(r[к['Выручка, руб']])
        сч['выручка известна'] += в > 0
        сч['≥1 млрд'] += в >= 1e9
        сч['100 млн – 1 млрд'] += 1e8 <= в < 1e9
        сч['50–100 млн'] += 5e7 <= в < 1e8
        сч['до 50 млн'] += 0 < в < 5e7
        сч['выручка неизвестна'] += в <= 0
        сч['с паспортом'] += any(r[n] for n in паспорт)
        сч['с ЛПР'] += bool(r[к['ЛПР: ФИО']])
        сч['получатель письма определён'] += bool(r[к['Письмо: e-mail']])
        сч['с мобильным'] += число(r[к['Мобильных']]) > 0
        сч['без ИНН (определены по сайту)'] += str(r[к['ИНН / УНП']] or '').startswith('САЙТ:')
    сегм = collections.Counter(str(r[к['Сегмент']] or '—') for r in отбор)

    out = openpyxl.Workbook()
    ws = out.active
    ws.title = 'Сводка'
    жир = Font(bold=True)
    ws.append(['База КЦ — компании с почтой (все, выручка любая; где известна — указана с источником)'])
    ws['A1'].font = жир
    ws.append(['Источник: %s, лист «Компании». Порядок: сначала с известной выручкой по убыванию, затем без выручки.' % вход.split('/')[-1]])
    ws.append([])
    ws.append(['Компаний в полном файле', всего])
    ws.append(['С почтой (в этом файле)', len(отбор)])
    for k in ('выручка известна', '≥1 млрд', '100 млн – 1 млрд', '50–100 млн', 'до 50 млн', 'выручка неизвестна',
              'с паспортом', 'с ЛПР', 'получатель письма определён', 'с мобильным', 'без ИНН (определены по сайту)'):
        ws.append(['  ' + k, сч[k]])
    ws.append([])
    ws.append(['Сегменты (топ-40)'])
    ws.cell(ws.max_row, 1).font = жир
    for s, n in сегм.most_common(40):
        ws.append(['  ' + s, n])
    ws.column_dimensions['A'].width = 70

    w2 = out.create_sheet('Компании')
    w2.append(шапка)
    for c in w2[1]:
        c.font = жир
    for r in отбор:
        w2.append(r)
    w2.freeze_panes = 'C2'
    for n in range(1, len(шапка) + 1):
        w2.column_dimensions[get_column_letter(n)].width = 18

    if 'Контакты' in wb.sheetnames:
        it = wb['Контакты'].iter_rows(values_only=True)
        ш = list(next(it))
        ki = ш.index('ИНН / УНП')
        w3 = out.create_sheet('Контакты')
        w3.append(ш)
        for c in w3[1]:
            c.font = жир
        for r in it:
            if r and str(r[ki]) in инн:
                w3.append(list(r))
        w3.freeze_panes = 'C2'
    out.save(выход)
    print('всего', всего, 'с почтой', len(отбор), dict(сч))


if __name__ == '__main__':
    main(*sys.argv[1:3])
