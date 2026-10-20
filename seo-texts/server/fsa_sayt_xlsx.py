# -*- coding: utf-8 -*-
"""Изготовители из деклараций ФСА, чьи контакты проверены по их собственному сайту (20.10, владелец: «дай отдельно тех,
чьи контакты ещё проверялись по самому сайту»): строки MEYER-7 «Компании»/«Контакты», где компания пришла из реестра
деклараций ФСА и сайт в базе есть (прошёл проверку «сайт той ли компании» и опровергателей). Плюс из деклараций: продукция,
число и дата деклараций, ссылка на декларацию, контакты заявителя; предварительная пометка «под что» по пояснению отсева
и флаг «найден только по основе "рис"» (подстрока: приспособления, присоединение, Рислинг — часто не наша продукция).

    python3 fsa_sayt_xlsx.py <MEYER-7.xlsx> <папка fsa: izg, otsev, kontakty2> <out.xlsx>
"""
import collections
import json
import os
import re
import sys

import openpyxl
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE as ILL
from openpyxl.styles import Font, PatternFill

M = re.compile(r'сортир|фотосепар|сепарац|рентген|инспекц|посторонн|отбраков|калибр|контрол\w*\s+(?:качеств|фасов|продук|упаков|'
               r'сырь|включ)|сыпуч', re.I)
K = re.compile(r'компрессор|сжат\w* воздух|пневм|азот|воздух', re.I)


def jl(п):
    for s in open(п, encoding='utf-8', errors='replace'):
        try:
            yield json.loads(s)
        except ValueError:
            pass


def main(м7, D, выход):
    ots = {y['инн']: y for y in jl(os.path.join(D, 'fsa-meyer-otsev.jsonl'))}
    к2 = {y['инн']: y for y in jl(os.path.join(D, 'fsa-meyer-kontakty2.jsonl'))}
    изг = {z['инн']: z for z in jl(os.path.join(D, 'fsa-meyer-izg.jsonl'))}
    wb = openpyxl.load_workbook(м7, read_only=True)
    it = wb['Компании'].iter_rows(values_only=True)
    ш = list(next(it))
    I = {k: i for i, k in enumerate(ш)}
    доп = ['Под что (предварительно, по пояснению отсева)', 'Пояснение отсева', 'Найдена только по основе «рис» ⚠',
           'Продукция (декларации, до 6)', 'Деклараций', 'Последняя декларация', 'Декларация (ссылка)',
           'Почты из декларации', 'Телефоны из декларации', 'Сайт из декларации']
    строки, инн = [], set()
    сч = collections.Counter()
    for r in it:
        if not r or not str(r[I['Откуда компания']] or '').startswith('реестр деклараций ФСА') or not r[I['Сайт']]:
            continue
        i = str(r[I['ИНН / УНП']])
        y, z, к = ots.get(i) or {}, изг.get(i) or {}, к2.get(i) or {}
        п = y.get('почему', '')
        под = 'Meyer (сортировка/рентген)' if M.search(п) else ('компрессор' if K.search(п) else 'неясно — смотреть продукцию')
        рис = (z.get('слова') or []) == ['рис']
        строки.append(list(r) + [под, п, 'да' if рис else '', ' | '.join(z.get('продукция') or []), z.get('деклараций'),
                                 z.get('последняя'), ('https://pub.fsa.gov.ru/rds/declaration/view/%s/common' % к['id']) if к.get('id') else '',
                                 ', '.join(к.get('почты') or []), ', '.join(к.get('телефоны') or []), ', '.join(к.get('сайты') or [])])
        инн.add(i)
        сч[под] += 1
        сч['только «рис»'] += рис
        сч['с почтой'] += bool(r[I['Почт']])
        сч['с номером'] += bool(r[I['Номеров']])
        сч['с ЛПР'] += bool(r[I['ЛПР: ФИО']] or r[I['ЛПР: телефон']])
        сч['с паспортом'] += any(r[I[k]] for k in ш if str(k).startswith('Паспорт'))
    порядок = {'Meyer (сортировка/рентген)': 0, 'неясно — смотреть продукцию': 1, 'компрессор': 2}
    строки.sort(key=lambda r: (порядок[r[len(ш)]], r[len(ш) + 2] == 'да', -(float(r[I['Выручка, руб']] or 0) if str(r[I['Выручка, руб']] or '').replace('.', '', 1).isdigit() else 0)))
    out = openpyxl.Workbook()
    sv = out.active
    sv.title = 'Сводка'
    sv.append(['Изготовители из деклараций ФСА с контактами, проверенными по собственному сайту (из MEYER-7, 20.10)'])
    sv['A1'].font = Font(bold=True)
    sv.append(['Сайт — свой (проверка сайта + опровергатели), контакты — с живых страниц сайта (лесенка Meyer), плюс контакты из '
               'декларации. «Под что» — предварительно по короткому пояснению модели отсева; точная классификация — отдельно.'])
    sv.append([])
    sv.append(['Компаний', len(строки)])
    for k in ('Meyer (сортировка/рентген)', 'неясно — смотреть продукцию', 'компрессор', 'только «рис»', 'с почтой', 'с номером',
              'с ЛПР', 'с паспортом'):
        sv.append(['  ' + k, сч[k]])
    sv.column_dimensions['A'].width = 70
    w = out.create_sheet('Компании')
    w.append(ш + доп)
    for c in w[1]:
        c.font = Font(bold=True)
    жёлтый = PatternFill('solid', fgColor='FFF2CC')
    for r in строки:
        w.append([ILL.sub(' ', v) if isinstance(v, str) else v for v in r])
        if r[len(ш) + 2] == 'да':
            for c in w[w.max_row]:
                c.fill = жёлтый
    w.freeze_panes = 'C2'
    w.auto_filter.ref = w.dimensions
    it = wb['Контакты'].iter_rows(values_only=True)
    шк = list(next(it))
    ki = шк.index('ИНН / УНП')
    wk = out.create_sheet('Контакты')
    wk.append(шк)
    for c in wk[1]:
        c.font = Font(bold=True)
    n = 0
    for r in it:
        if r and str(r[ki]) in инн:
            wk.append([ILL.sub(' ', v) if isinstance(v, str) else v for v in r])
            n += 1
    wk.freeze_panes = 'C2'
    out.save(выход)
    print(len(строки), dict(сч), 'контактов', n)


if __name__ == '__main__':
    main(*sys.argv[1:4])
