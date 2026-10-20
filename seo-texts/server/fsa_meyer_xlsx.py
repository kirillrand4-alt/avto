# -*- coding: utf-8 -*-
"""FSA-CERBER-MEYER.xlsx — изготовители из деклараций ФСА + реестр Цербера (18.10; 20.10 — итоговая сборка).

20.10, добавлено: карточки checko по юрлицам (статус, основной и доп. ОКВЭД, выручка с годом, руководитель, регион,
число госзакупок/проверок, ссылка на карточку), DaData — где checko нет; сайт и контакты из повторного прохода
карточек деклараций (только блоки заявителя/изготовителя, без органа сертификации) со ссылкой на декларацию;
ИНН Цербера — сверенные по самому реестру; отметка «прошёл конвейер Meyer-7» и активные сделки Meyer из свежей
выгрузки Битрикса. Номера checko не берём (владелец 19.10). Ликвидированные — отдельный лист.

    python3 fsa_meyer_xlsx.py <папка данных> <meyer7-spisok.json> <meyer7-sdelki-aktivnye.tsv> <out.xlsx>
"""
import collections
import json
import os
import re
import sys

from openpyxl import Workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE as ILL
from openpyxl.styles import Font, PatternFill

чист = lambda r: [ILL.sub(' ', v) if isinstance(v, str) else v for v in r]  # noqa: E731
ПОР = 5  # почта/телефон у 5+ изготовителей — оформитель деклараций, не контакт компании
БЕСПЛ = {'mail.ru', 'bk.ru', 'inbox.ru', 'list.ru', 'yandex.ru', 'ya.ru', 'gmail.com', 'rambler.ru', 'yahoo.com', 'icloud.com',
         'outlook.com', 'hotmail.com', 'internet.ru', 'yandex.com', 'mail.com', 'lenta.ru', 'autorambler.ru', 'ro.ru',
         'myrambler.ru', 'km.ru', 'narod.ru', 'yandex.by', 'tut.by'}
МЁРТВ = re.compile(r'ликвид|прекрат|банкрот|LIQUIDAT|BANKRUPT|недейств', re.I)


def f(v):
    try:
        return float(v or 0)
    except (TypeError, ValueError):
        return 0.0


def jl(путь):
    if not os.path.exists(путь):
        return
    for s in open(путь, encoding='utf-8', errors='replace'):
        try:
            yield json.loads(s)
        except ValueError:
            pass


def моб(t):
    return re.sub(r'\D', '', t)[-10:].startswith('9')


def main(D, путь_списка, путь_сделок, выход):
    ots = {y['инн']: y for y in jl(os.path.join(D, 'fsa-meyer-otsev.jsonl'))}
    к2 = {y['инн']: y for y in jl(os.path.join(D, 'fsa-meyer-kontakty2.jsonl'))}
    чк = {}
    for y in jl(os.path.join(D, 'fsa-meyer-checko.jsonl')):
        if y.get('итог') == 'ok' and len(y.get('inn', '')) == 10:
            чк[y['inn']] = y
    дд = {y['инн']: y for y in jl(os.path.join(D, 'fsa-meyer-dadata.jsonl'))}
    м7 = set(json.load(open(путь_списка, encoding='utf-8'))['компании'])
    сд = {}
    for s in open(путь_сделок, encoding='utf-8'):
        ч = s.rstrip('\n').split('\t')
        if len(ч) > 2 and ч[2] == 'ИНН':
            сд[ч[0]] = ч[1]
    изг = list(jl(os.path.join(D, 'fsa-meyer-izg.jsonl')))
    кп, кт = collections.Counter(), collections.Counter()
    for z in изг:
        к = к2.get(z['инн']) or {}
        for a in set(к.get('почты') or z['почты']):
            кп[a] += 1
        for t in set(к.get('телефоны') or z['телефоны']):
            кт[re.sub(r'\D', '', t)[-10:]] += 1

    wb = Workbook()
    sv = wb.active
    sv.title = 'Сводка'
    H = ['ИНН', 'Изготовитель', 'Статус', 'Регион', 'Адрес (из декларации)', 'ОКВЭД основной', 'ОКВЭД доп. (с карточки)',
         'Выручка, руб', 'Год выручки', 'Откуда выручка', 'Руководитель', 'Деклараций', 'Последняя декларация',
         'Продукция (до 6)', 'Почты (из декларации)', 'Мобильные', 'Городские', 'Посредник? (контакт у 5+ компаний)',
         'Сайт (из декларации)', 'Сайт (по домену почты)', 'Декларация (ссылка)', 'Карточка checko', 'Госзакупок (checko)',
         'Проверок (checko)', 'Уже в базе Meyer', 'Прошёл конвейер Meyer-7 (контакты с сайта — в MEYER-7.xlsx)',
         'Активная сделка Meyer', 'Оценка модели', 'Почему']
    w1, w2, w4 = wb.create_sheet('Изготовители'), wb.create_sheet('Отсеяно'), wb.create_sheet('Ликвидированы')
    for w in (w1, w2, w4):
        w.append(H)
        for c in w[1]:
            c.font = Font(bold=True)
            c.fill = PatternFill('solid', fgColor='DDEBF7')
    сч = collections.Counter()
    rows = []
    for z in изг:
        инн = z['инн']
        y = ots.get(инн) or {}
        отв = {'da': 'да', '可能': 'возможно', '是': 'да'}.get(y.get('ответ', ''), y.get('ответ', ''))
        к = к2.get(инн) or {}
        с = чк.get(инн) or {}
        д = дд.get(инн) or {}
        почты_все = к.get('почты') if к else z['почты']
        тел_все = к.get('телефоны') if к else z['телефоны']
        почты = [a for a in почты_все or [] if кп[a] < ПОР]
        тел = [t for t in тел_все or [] if кт[re.sub(r'\D', '', t)[-10:]] < ПОР]
        поср = [a for a in почты_все or [] if кп[a] >= ПОР] + [t for t in тел_все or [] if кт[re.sub(r'\D', '', t)[-10:]] >= ПОР]
        статус = с.get('статус') or {'ACTIVE': 'Действующая', 'LIQUIDATED': 'Ликвидирована', 'LIQUIDATING': 'Ликвидируется',
                                     'BANKRUPT': 'Банкротство', 'REORGANIZING': 'Реорганизация'}.get(д.get('статус'), д.get('статус') or '')
        выр, год, откуда = f(с.get('выручка')), с.get('выручка_год') or '', 'checko' if с.get('выручка') else ''
        if not выр and f(z.get('выручка')):
            выр, год, откуда = f(z['выручка']), '', 'база обзвона'
        оквэд = с.get('оквэд') or д.get('оквэд') or z.get('оквэд') or ''
        доп = [x for x in (с.get('оквэд_все') or []) if x != оквэд]
        рук = ' — '.join(x for x in ((с.get('руководитель') or д.get('руководитель') or ''), (с.get('должность') or д.get('должность') or '')) if x)
        регион = с.get('регион') or д.get('регион') or z.get('регион') or ''
        сайт_д = ', '.join(к.get('сайты') or [])
        сайт_п = ', '.join(sorted({a.split('@')[-1] for a in почты if a.split('@')[-1] not in БЕСПЛ}))
        ссылка = ('https://pub.fsa.gov.ru/rds/declaration/view/%s/common' % к['id']) if к.get('id') else ''
        сделка = сд.get(инн, '')
        r = [инн, z['имя'], статус, регион, z['адрес'], оквэд, ', '.join(доп[:15]), выр or None, год, откуда, рук,
             z['деклараций'], z['последняя'], ' | '.join(z['продукция']), ', '.join(почты), ', '.join(t for t in тел if моб(t)),
             ', '.join(t for t in тел if not моб(t)), ', '.join(поср), сайт_д, сайт_п, ссылка, с.get('ссылка', ''),
             с.get('госзакупки') or None, с.get('проверки') or None, z['в_базе_meyer'], 'да' if инн in м7 else '', сделка, отв,
             y.get('почему', '')]
        куда = 'ликв' if МЁРТВ.search(статус or '') else ('отсев' if отв == 'нет' else 'список')
        rows.append((куда, выр, f(z['деклараций']), r, почты, тел, сайт_д or сайт_п, инн))
    rows.sort(key=lambda t: (-t[1], -t[2]))
    for куда, выр, _, r, почты, тел, сайт, инн in rows:
        {'ликв': w4, 'отсев': w2, 'список': w1}[куда].append(чист(r))
        сч[куда] += 1
        if куда == 'список':
            сч['с почтой'] += bool(почты)
            сч['с телефоном'] += bool(тел)
            сч['с мобильным'] += any(моб(t) for t in тел)
            сч['с сайтом'] += bool(сайт)
            сч['юрлица'] += len(инн) == 10
            сч['checko-карточка'] += инн in чк
            сч['выручка известна'] += выр > 0
            сч['выручка от 100 млн'] += выр >= 1e8
            сч['прошёл конвейер Meyer-7'] += инн in м7
            сч['активная сделка Meyer'] += bool(сд.get(инн))
            сч['новые для Meyer'] += not r[24]
    # Цербер: ИНН, сверенный по самому реестру (поиск реестра по ИНН выдаёт это юрлицо)
    w3 = wb.create_sheet('Цербер')
    w3.append(['Номер в реестре', 'Вид предприятия', 'Объект', 'Адрес объекта', 'Юрлицо', 'Юр. адрес', 'ИНН', 'Как определён ИНН'])
    for c in w3[1]:
        c.font = Font(bold=True)
    инн2 = {(y['юрлицо'], y['рег']): y for y in jl(os.path.join(D, 'cerber-meyer-inn2.jsonl'))}

    def рег(a):
        m = re.search(r'Российская Федерация, ([^,]+)', a or '')
        return (m.group(1) if m else '').replace('обл.', '').replace('край', '').replace('Респ.', '').strip().lower()[:12]
    nc = 0
    for x in jl(os.path.join(D, 'cerber-meyer.jsonl')):
        if not x.get('номер'):
            continue
        y = инн2.get((x['юрлицо'], рег(x.get('юр_адрес') or x.get('адрес')))) or {}
        w3.append(чист([x['номер'], x['вид'], x['объект'], x['адрес'], x['юрлицо'], x['юр_адрес'], y.get('инн', ''),
                        y.get('как', '')]))
        nc += 1
        сч['Цербер: ИНН подтверждён реестром'] += bool(y.get('инн'))
    sv.append(['ФСА (декларации соответствия, 238 основ продукции Meyer) + Цербер (реестр предприятий ТС). Сборка 20.10']); sv['A1'].font = Font(bold=True)
    sv.append(['Данные компаний: checko (юрлица), DaData; контакты — из карточек деклараций (заявитель = изготовитель), '
               'посредники (контакт у 5+ компаний) вынесены; номера checko не используются.'])
    sv.append([])
    for k in ('список', 'отсев', 'ликв'):
        sv.append([{'список': 'Изготовители (в работе)', 'отсев': 'Отсеяно (модель/правило)', 'ликв': 'Ликвидированы/банкроты'}[k], сч[k]])
    for k in ('юрлица', 'checko-карточка', 'выручка известна', 'выручка от 100 млн', 'с почтой', 'с телефоном', 'с мобильным',
              'с сайтом', 'новые для Meyer', 'прошёл конвейер Meyer-7', 'активная сделка Meyer'):
        sv.append(['  ' + k, сч[k]])
    sv.append(['Цербер: объектов', nc])
    sv.append(['Цербер: ИНН подтверждён реестром', сч['Цербер: ИНН подтверждён реестром']])
    sv.column_dimensions['A'].width = 60
    for w in (w1, w2, w4):
        w.freeze_panes = 'C2'
        w.auto_filter.ref = w.dimensions
        for n, wd in enumerate((13, 40, 16, 22, 40, 10, 30, 14, 8, 12, 30, 10, 12, 60, 30, 18, 18, 24, 22, 22, 30, 30, 10, 10, 12,
                                14, 30, 10, 30), 1):
            w.column_dimensions[w.cell(row=1, column=n).column_letter].width = wd
    wb.save(выход)
    print(dict(сч), nc)


if __name__ == '__main__':
    main(*sys.argv[1:5])
