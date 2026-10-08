# -*- coding: utf-8 -*-
"""Колонка «Есть контакт в Битрикс» во всех листах Excel, где есть «ИНН» (владелец 08.10).

Источник — выгрузка сделок Битрикса (CSV: ID, TITLE, ВОРОНКА, INN). Совпадение по ИНН ->
«да: N сделок (воронки …)», иначе пусто. Колонка встаёт сразу после «ИНН».

    python dobavit_bitrix.py <deals_inn.csv> <in.xlsx> <out.xlsx>
"""
import collections
import csv
import json
import re
import sys

from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill


def сделки(п):
    по_инн = collections.defaultdict(list)
    with open(п, encoding='utf-8-sig', newline='') as f:
        for r in csv.DictReader(f):
            i = re.sub(r'\D', '', r.get('INN') or '')
            if len(i) in (10, 12):
                по_инн[i].append(r.get('ВОРОНКА') or '')
    return по_инн


def main(п_csv, п_in, п_out):
    по_инн = сделки(п_csv)
    wb = load_workbook(п_in)
    сч = {}
    for ws in wb.worksheets:
        шапка = [c.value for c in ws[1]]
        if 'ИНН' not in шапка:
            continue
        к = шапка.index('ИНН') + 1
        ws.insert_cols(к + 1)
        h = ws.cell(row=1, column=к + 1, value='Есть контакт в Битрикс')
        h.font = Font(bold=True, color='FFFFFF')
        h.fill = PatternFill('solid', fgColor='305496')
        h.alignment = Alignment(wrap_text=True, vertical='top')
        ws.column_dimensions[h.column_letter].width = 26
        n = 0
        for r in range(2, ws.max_row + 1):
            i = re.sub(r'\D', '', str(ws.cell(row=r, column=к).value or ''))
            вв = по_инн.get(i)
            if вв:
                n += 1
                в = collections.Counter(вв).most_common(3)
                ws.cell(row=r, column=к + 1, value='да: %d сдел. (%s)' % (len(вв), ', '.join('%s ×%d' % x for x in в)))
        if ws.auto_filter.ref:
            ws.auto_filter.ref = 'A1:%s%d' % (ws.cell(row=1, column=ws.max_column).column_letter, ws.max_row)
        сч[ws.title] = '%d из %d строк' % (n, ws.max_row - 1)
    wb.save(п_out)
    print(json.dumps({'ИНН в выгрузке': len(по_инн), 'совпало по листам': сч}, ensure_ascii=False, indent=1))


if __name__ == '__main__':
    main(*sys.argv[1:4])
