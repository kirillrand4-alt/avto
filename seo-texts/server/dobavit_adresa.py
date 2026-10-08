# -*- coding: utf-8 -*-
"""Лист «Компании»: колонка «Где найдены контакты» сразу после «Сайт» (владелец 08.10) — адреса страниц,
с которых взяты номера этой компании (лист «Контакты», «Ссылка на источник»), без повторов; сначала
страницы с номерами покупателей оборудования/ЛПР, потом остальные. Каждый адрес с новой строки.

    python dobavit_adresa.py <in.xlsx> <out.xlsx>   (можно in == out)
"""
import collections
import sys

from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill

ПОКУПАТЕЛИ = ('закупки/снабжение', 'главный механик', 'главный энергетик', 'главный инженер', 'технический директор',
              'производство')


def main(п_in, п_out):
    wb = load_workbook(п_in)
    к = wb['Контакты']
    шк = [c.value for c in к[1]]
    иИНН, иСс, иРоль = (шк.index(x) for x in ('ИНН', 'Ссылка на источник', 'Роль'))
    иЛПР = шк.index('ЛПР КЦ') if 'ЛПР КЦ' in шк else шк.index('ЛПР Meyer')
    адреса = collections.defaultdict(list)  # ИНН -> [(приоритет, url)]
    for r in к.iter_rows(min_row=2, values_only=True):
        u = (r[иСс] or '').strip()
        if not u:
            continue
        пр = 0 if r[иРоль] in ПОКУПАТЕЛИ else 1 if r[иЛПР] else 2
        адреса[str(r[иИНН])].append((пр, u))
    ws = wb['Компании']
    шапка = [c.value for c in ws[1]]
    if 'Где найдены контакты' in шапка:
        кол = шапка.index('Где найдены контакты') + 1
    else:
        кол = шапка.index('Сайт') + 2
        ws.insert_cols(кол)
        h = ws.cell(row=1, column=кол, value='Где найдены контакты')
        h.font = Font(bold=True, color='FFFFFF')
        h.fill = PatternFill('solid', fgColor='305496')
        h.alignment = Alignment(wrap_text=True, vertical='top')
        ws.column_dimensions[h.column_letter].width = 55
    иИННк = шапка.index('ИНН') + 1
    n = 0
    for row in range(2, ws.max_row + 1):
        i = str(ws.cell(row=row, column=иИННк).value)
        лучшие = {}
        for пр, u in адреса.get(i, []):
            лучшие[u] = min(пр, лучшие.get(u, 9))
        список = [u for u, _ in sorted(лучшие.items(), key=lambda x: x[1])]
        if список:
            n += 1
            c = ws.cell(row=row, column=кол, value='\n'.join(список))
            c.alignment = Alignment(wrap_text=False, vertical='top')
    if ws.auto_filter.ref:
        ws.auto_filter.ref = 'A1:%s%d' % (ws.cell(row=1, column=ws.max_column).column_letter, ws.max_row)
    wb.save(п_out)
    print(п_out, 'компаний с адресами:', n, 'из', ws.max_row - 1)


if __name__ == '__main__':
    main(*sys.argv[1:3])
