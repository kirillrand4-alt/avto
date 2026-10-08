# -*- coding: utf-8 -*-
"""Файл 6 Meyer (08.10): сборка из набора meyer6 тем же сборщиком (kc_xlsx, режим Meyer) + оформление под Meyer:
колонки «ЛПР Meyer», «Раздел» (основной список КЦ / ниже порога), сводка про файл 6, адреса страниц-источников,
отметка Битрикса.

    python meyer6_xlsx.py <папка с meyer6-*.json(l)> <папка файлов Meyer 1–5> <deals_inn.csv> <out.xlsx>
"""
import json
import os
import sys
import time

from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill

import dobavit_adresa
import dobavit_bitrix
import kc_xlsx


def main(п, п_meyer, п_csv, п_out):
    ф = lambda x: os.path.join(п, 'meyer6-' + x)  # noqa: E731
    kc_xlsx.main(ф('spisok.json'), ф('kontakty.jsonl'), п_meyer, п_out, None, ф('audit.jsonl'), ф('audit2.jsonl'),
                 ф('sayt-proverka.jsonl'), ф('glubokiy2.jsonl'))
    сп = json.load(open(ф('spisok.json'), encoding='utf-8'))['компании']
    wb = load_workbook(п_out)
    for ws in wb.worksheets:
        шапка = [c.value for c in ws[1]]
        for j, h in enumerate(шапка):
            if h in ('ЛПР КЦ', 'Из них ЛПР КЦ'):
                ws.cell(row=1, column=j + 1, value=h.replace('КЦ', 'Meyer'))
        if 'Из них покупатели оборудования' in шапка:
            ws.delete_cols(шапка.index('Из них покупатели оборудования') + 1)
        if ws.title == 'Компании':
            шапка = [c.value for c in ws[1]]
            к = шапка.index('Сегмент') + 2
            ws.insert_cols(к)
            h = ws.cell(row=1, column=к, value='Раздел')
            h.font = Font(bold=True, color='FFFFFF')
            h.fill = PatternFill('solid', fgColor='305496')
            h.alignment = Alignment(wrap_text=True, vertical='top')
            ws.column_dimensions[h.column_letter].width = 22
            ки = шапка.index('ИНН') + 1
            for r in range(2, ws.max_row + 1):
                i = str(ws.cell(row=r, column=ки).value)
                ws.cell(row=r, column=к, value=(сп.get(i) or {}).get('раздел_кц', ''))
    св = wb['Сводка']
    for r in range(1, св.max_row + 1):
        v = св.cell(row=r, column=1).value or ''
        if v.startswith('База для КЦ'):
            св.cell(row=r, column=1, value='Файл 6 Meyer — %s. Компании из сбора поиском (молоко, сыры, мясо, хлеб, корма; '
                                         'напитки исключены): выручка от 30 млн или неизвестна, выше и ниже 1,5 млрд. '
                                         'Компании из файлов Meyer 1–5 не убирались — отмечены в колонке «Есть в файлах Meyer».'
                                         % time.strftime('%d.%m.%Y'))
        elif v.startswith('Порядок:'):
            св.cell(row=r, column=1, value='Порядок контактов: ЛПР Meyer (руководитель, главный инженер, техдиректор, '
                                         'производство, главный технолог/технолог, качество, закупки) с ФИО -> без ФИО -> '
                                         'техслужбы -> прочие -> общие/приёмные; мобильные выше.')
        elif 'ЛПР КЦ' in v or 'покупателя оборудования' in v:
            св.cell(row=r, column=1, value=v.replace('ЛПР КЦ', 'ЛПР Meyer'))
            if 'покупателя оборудования' in v:
                св.cell(row=r, column=1, value=None)
                св.cell(row=r, column=2, value=None)
    wb.save(п_out)
    dobavit_adresa.main(п_out, п_out)
    dobavit_bitrix.main(п_csv, п_out, п_out)


if __name__ == '__main__':
    main(*sys.argv[1:5])
