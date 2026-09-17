# -*- coding: utf-8 -*-
"""Эксель по заявителям технологического присоединения из СиПР Системного оператора.

Что это за источник. Обосновывающие материалы схемы и программы развития энергосистем
(СиПР СО ЕЭС 2025-2030), таблица «Перечень планируемых к присоединению потребителей».
Там названы: кто присоединяется, к какому центру питания, СКОЛЬКО МВт добавляет и В КАКОМ
ГОДУ. Это ранняя стадия с датой и масштабом — то, чего лента новостей не даёт.

Листы:
  «Заявители»     все 301 строка, отсортированы по величине прироста мощности
  «Наша тема»     30 строк, где проект похож на нашу тему по признакам
                  (воздухоразделение, газопереработка, СПГ, металлургия, цемент, целлюлоза)
  «Как читать»    откуда взято, какие контроли стояли и чего в файле НЕТ

Счётчики и разметка считаются в питоне: это снимок на дату, а не модель. В прошлой
эксельке формулами было 2 475 COUNTIFS, и LibreOffice не пересчитал их ни за 175 секунд,
ни за 600 — у владельца файл открывался бы так же долго.
"""
import csv
import io
import os

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

KAT = os.path.dirname(os.path.abspath(__file__))
VHOD = os.path.join(KAT, 'SIPR-ZAYAVITELI-INN.csv')
VYHOD = os.path.join(KAT, 'SIPR-ZAYAVITELI-TP.xlsx')
SHRIFT = 'Arial'

r = list(csv.DictReader(io.open(VHOD, encoding='utf-8-sig'), delimiter=';'))
print('строк на входе: %d' % len(r))


def chislo(x, k):
    t = (x.get(k) or '').replace(',', '.').strip()
    try:
        return float(t)
    except Exception:  # noqa: BLE001
        return None


def flag(x, k):
    v = (x.get(k) or '').strip()
    return '' if v in ('', '0', 'нет') else 'да'


r.sort(key=lambda x: chislo(x, 'uvelichenie_MVt_chislo') or -1, reverse=True)

KOL = [
    ('Заявитель', 'naimenovanie'), ('Регион', 'region'),
    ('Что за проект', 'proekt'),
    ('Прирост мощности, МВт', 'uvelichenie_MVt_chislo'),
    ('Было присоединено, МВт', 'ranee_MVt_chislo'),
    ('Год ввода', 'god_chislo'), ('Напряжение, кВ', 'napryazhenie_kV'),
    ('Центр питания', 'centr_pitaniya'),
    ('ИНН', 'inn'), ('ОГРН', 'ogrn'), ('Статус в ЕГРЮЛ', 'status_egryul'),
    ('Руководитель', 'rukovoditel'), ('Адрес по ЕГРЮЛ', 'adres_egryul'),
    ('Как найден ИНН', 'razreshenie'), ('Новое для нашей базы', 'novoe_dlya_bazy'),
    ('Кандидатов на ИНН', 'kandidatov'), ('Кандидаты (ИНН)', 'kandidaty_inn'),
    ('Похоже на нашу тему', 'priznak_centrobezhnogo'),
    ('Заявитель сомнителен', 'zayavitel_somnitelen'),
    ('Мощность сомнительна', 'moshchnost_somnitelna'),
    ('Файл-источник', 'fajl'), ('Страница', 'stranica'),
]

wb = Workbook()

# ------------------------------------------------------------------ Как читать
ws = wb.active
ws.title = 'Как читать'
Z = Font(name=SHRIFT, size=12, bold=True)
O = Font(name=SHRIFT, size=10)
nasha = [x for x in r if (x.get('priznak_centrobezhnogo') or '').strip()]
s_inn = [x for x in r if (x.get('inn') or '').strip()]
novye = [x for x in r if (x.get('novoe_dlya_bazy') or '').strip().lower() == 'да']
uzhe_est = [x for x in r if (x.get('novoe_dlya_bazy') or '').strip().lower() == 'нет']
somn_z = [x for x in r if flag(x, 'zayavitel_somnitelen')]
somn_m = [x for x in r if flag(x, 'moshchnost_somnitelna')]
OPIS = [
    ('Заявители технологического присоединения: кто наращивает мощность, сколько и когда', True),
    ('', False),
    ('ИСТОЧНИК. Обосновывающие материалы СиПР СО ЕЭС 2025-2030, таблица «Перечень', False),
    ('   планируемых к присоединению потребителей». По файлу на субъект РФ. Там названы', False),
    ('   заявитель, центр питания, прирост мощности в МВт и год ввода — то есть событие', False),
    ('   ранней стадии, у которого сразу есть дата и масштаб.', False),
    ('   Почему это ценно: лента новостей приносит событие, когда стройка уже видна. Здесь', False),
    ('   предприятие названо на этапе, когда оно только запрашивает мощность.', False),
    ('', False),
    ('ЧИСЛА ЭТОГО ФАЙЛА', True),
    ('   строк ................................. %d  (74 региона)' % len(r), False),
    ('   с приростом мощности .................. %d' % len([x for x in r if (chislo(x, 'uvelichenie_MVt_chislo') or 0) > 0]), False),
    ('   с годом ввода ......................... %d' % len([x for x in r if (x.get('god_chislo') or '').strip()]), False),
    ('   ИНН найден однозначно ................. %d  (разных предприятий %d)' % (len(s_inn), len({x['inn'] for x in s_inn})), False),
    ('   из них НОВЫХ для нашей базы ........... %d' % len(novye), False),
    ('   уже были в базе ....................... %d' % len(uzhe_est), False),
    ('   проверить не на чем (нет ИНН) ......... %d' % len([x for x in r if not (x.get('inn') or '').strip()]), False),
    ('   похоже на нашу тему ................... %d  (лист «Наша тема»)' % len(nasha), False),
    ('', False),
    ('ЧЕГО В ФАЙЛЕ НЕТ И ЧТО В НЁМ СОМНИТЕЛЬНО, сказано прямо:', True),
    ('   • 301 строка — это НИЖНЯЯ граница. У 24 регионов часть строк не снялась, у 12', False),
    ('     таблица не найдена вовсе. В отчёте по каждому записано «взято N из примерно M».', False),
    ('   • %d строк помечены «заявитель сомнителен»: в исходной ячейке слиплись две строки,' % len(somn_z), False),
    ('     и имя может быть обрезано или склеено с соседним. Не выброшены — имя читается,', False),
    ('     а колонка «Страница» позволяет проверить по первоисточнику за минуту.', False),
    ('   • %d строк помечены «мощность сомнительна».' % len(somn_m), False),
    ('   • людей в источнике нет. Руководитель, где он есть, взят из ЕГРЮЛ, а не из СиПР.', False),
    ('   • у 17 строк ИНН НЕ проставлен, хотя кандидаты есть: регион их не различает.', False),
    ('     Они лежат в колонке «Кандидаты (ИНН)» — лучше пусто, чем сшить не с тем юрлицом.', False),
    ('', False),
    ('ПОПРАВКА, КОТОРУЮ НАДО ЗНАТЬ', True),
    ('   В первом заходе было заявлено 525 заявителей. Это оказался пересчёт: тот замер брал', False),
    ('   все юрлица в 8 000 знаках после заголовка, куда попадает и текст прогноза', False),
    ('   потребления. Сверка по регионам: Вологодская — было 5, реально 1 (Северсталь,', False),
    ('   +86 МВт); Кабардино-Балкария — 10 против 3. Честное число построчного съёма — 301.', False),
    ('', False),
    ('КОНТРОЛИ', True),
    ('   • таблица снимается ПО КООРДИНАТАМ (x задаёт колонку, y строку), а не текстом:', False),
    ('     обычное извлечение отдаёт ячейки колонка-за-колонкой и порядок строк теряется.', False),
    ('   • положительный контроль добора ИНН: ММК обязан находиться (ИНН 7414003633).', False),
    ('     Первая версия дописывала регион в строку запроса и давала 0 карточек даже по ММК —', False),
    ('     отрицательный контроль этого не заметил бы, он честно показывал ноль.', False),
    ('   • пять дефектов извлечения поймано и исправлено, каждый ронял строки МОЛЧА:', False),
    ('     регистрозависимая шапка (Пермский край давал 0 строк при пяти заявителях, включая', False),
    ('     Уралкалий +58,5 МВт), колонцифры как ложные якоря строк, хвост шапки в имени,', False),
    ('     слишком широкая граница шапки срезала самые крупные строки, номер строки слипался', False),
    ('     с названием проекта.', False),
]
for i, (t, zh) in enumerate(OPIS, start=1):
    c = ws.cell(row=i, column=1, value=t)
    c.font = Z if zh else O
ws.column_dimensions['A'].width = 103


def list_s_dannymi(imya, stroki):
    w = wb.create_sheet(imya)
    w.append([z for z, _ in KOL])
    for x in stroki:
        ryad = []
        for _, k in KOL:
            if k in ('uvelichenie_MVt_chislo', 'ranee_MVt_chislo'):
                ryad.append(chislo(x, k))
            elif k in ('zayavitel_somnitelen', 'moshchnost_somnitelna'):
                ryad.append(flag(x, k))
            elif k == 'god_chislo':
                ryad.append((x.get(k) or '').strip())
            else:
                ryad.append((x.get(k) or '').strip())
        w.append(ryad)
    return w


w1 = list_s_dannymi('Заявители', r)
w2 = list_s_dannymi('Наша тема', nasha)

SH_F = Font(name=SHRIFT, size=10, bold=True, color='FFFFFF')
SH_FILL = PatternFill('solid', fgColor='2F5496')
SHIR = [40, 22, 52, 13, 13, 9, 11, 30, 13, 15, 14, 26, 38, 22, 20, 10, 24, 16, 12, 12, 26, 9]
for w in (w1, w2):
    for c in w[1]:
        c.font = SH_F
        c.fill = SH_FILL
        c.alignment = Alignment(vertical='center', wrap_text=True)
    w.freeze_panes = 'A2'
    w.auto_filter.ref = w.dimensions
    w.row_dimensions[1].height = 32
    for i, sh in enumerate(SHIR, start=1):
        w.column_dimensions[get_column_letter(i)].width = sh
    for row in w.iter_rows(min_row=2):
        for c in row:
            c.font = Font(name=SHRIFT, size=10)
            c.alignment = Alignment(vertical='top')
        for j in (4, 5):
            row[j - 1].number_format = '#,##0.0'

wb.save(VYHOD)
print('файл: %s' % VYHOD)
print('  Заявители: %d строк · Наша тема: %d строк' % (len(r), len(nasha)))
print('  ИНН у %d, новых для базы %d, сомнительных заявителей %d' % (len(s_inn), len(novye), len(somn_z)))
